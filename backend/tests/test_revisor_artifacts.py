"""Revisor-facing artifacts — what the revisor actually opens and receives.

These tests READ THE CONTENT of the generated files (openpyxl for the xlsx, the
csv module for the CSV, pypdf for the PDFs) — a size or magic-bytes check is no
proof that a total is right. Every send is stubbed: no test reaches Resend.

Stage 1 — money & safety:
  · Excel totals exclude drafts and equal the Oversigt sheet, the range PDF and
    the e-mail body for the same period (one source of truth: period_totals).
  · The default CSV parses with ";" and a decimal comma in Danish Excel.
  · A staff-typed "<b>" Lukket af is escaped in the revisor mail HTML.
  · send-to-accountant only mails the SAVED revisor address and has a daily cap.
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import date, datetime, timedelta
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app import models as _all_models  # noqa: F401 — register all models
from app.main import app, _db_ready
from app.models.audit_log import AuditLog
from app.models.business_profile import BusinessProfile
from app.models.daily_close import DailyClose, encode_breakdown
from app.models.user import User
from app.services.auth import create_access_token, hash_password
from app.services.daily_close_range_export import (
    build_daily_close_range_pdf,
    build_daily_close_range_xlsx,
    closes_to_csv_bytes,
    period_totals,
)
from app.utils.time import utc_now

_db_ready.set()


# ─── Content readers ──────────────────────────────────────────────────


def pdf_text(pdf_bytes: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def _changes(html: str) -> list[str]:
    """The correction's changed figures, one per table row: 'Kort 10.000,00
    kr. → 4.000,00 kr.' — the mail renders them as Linje | Før | Nu."""
    out = []
    for row in re.findall(r"<tr>(.*?)</tr>", html, flags=re.S):
        cells = [re.sub(r"<[^>]+>", "", c) for c in re.findall(r"<td[^>]*>(.*?)</td>", row, flags=re.S)]
        if len(cells) == 3 and cells[1].endswith("kr.") and cells[2].endswith("kr."):
            out.append(f"{cells[0]} {cells[1]} → {cells[2]}")
    return out


def _eval_sumifs(ws, formula: str) -> float:
    """Evaluate the one formula shape the workbook writes:
    =SUMIFS(D2:D9,$T$2:$T$9,"Låst") — sum the column where Status matches."""
    m = re.fullmatch(
        r'=SUMIFS\(([A-Z]+)(\d+):([A-Z]+)(\d+),\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+),"([^"]+)"\)',
        formula,
    )
    assert m, f"unexpected totals formula: {formula}"
    col, r1, _, r2, scol, _s1, _, _s2, crit = m.groups()
    total = 0.0
    for r in range(int(r1), int(r2) + 1):
        if ws[f"{scol}{r}"].value == crit:
            v = ws[f"{col}{r}"].value
            if isinstance(v, (int, float)):
                total += v
    return round(total, 2)


# ─── Builders ─────────────────────────────────────────────────────────


def _close(d, rev, moms, *, status="confirmed", pay=None, cats=None, **kw):
    base = dict(
        id=uuid.uuid4(), user_id=uuid.uuid4(), branch_id=None, date=d,
        revenue_categories=encode_breakdown(cats if cats is not None else {"food": rev}),
        revenue_total=rev,
        payment_categories=encode_breakdown(pay if pay is not None else {"cash": rev}),
        payment_total=rev, moms_total=moms,
        revenue_ex_moms=(round(rev - moms, 2) if moms is not None else None),
        moms_mode="auto",
        cash_expected=(pay or {"cash": rev}).get("cash"),
        cash_counted=(pay or {"cash": rev}).get("cash"),
        cash_difference=0.0, tips_total=None, tips_staff_count=None,
        tips_per_person=None, status=status, notes=None, closed_by="Lars",
        closed_at=datetime(2026, 9, 25, 23, 28) if status == "confirmed" else None,
        is_deleted=False,
    )
    base.update(kw)
    return DailyClose(**base)


def _sep_period():
    """Two locked days and a draft day with a big figure — the shape of the
    rating's September export (a 30.053 kr. draft summed into 'bekræftede')."""
    return [
        _close(date(2026, 9, 24), 1000.0, 200.0),
        _close(date(2026, 9, 25), 2500.0, 500.0,
               pay={"cash": 500.0, "card": 1500.0, "gift_card": 500.0}),
        _close(date(2026, 9, 26), 30053.0, 6010.6, status="draft"),
    ]


# ─── 1. Excel totals exclude drafts and equal every other artifact ────


def test_xlsx_totals_row_excludes_drafts_and_equals_oversigt_pdf_and_mail():
    from openpyxl import load_workbook
    from app.routers.daily_close import _accountant_email_body

    closes = _sep_period()
    f, t = date(2026, 9, 1), date(2026, 9, 30)
    wb = load_workbook(io.BytesIO(build_daily_close_range_xlsx(
        closes, from_date=f, to_date=t, business_name="Mirabelle ApS", currency="DKK")))

    ws = wb["Kasserapport"]
    header = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    assert header[19] == "Status"
    # The draft is LISTED, and marked.
    statuses = [ws.cell(row=r, column=20).value for r in range(2, 5)]
    assert statuses == ["Låst", "Låst", "Kladde (ikke medregnet)"]
    totals_row = 5
    assert ws.cell(row=totals_row, column=1).value == "I alt — 2 låste dage"
    rev_formula = ws.cell(row=totals_row, column=4).value
    moms_formula = ws.cell(row=totals_row, column=5).value
    assert "SUMIFS" in rev_formula and "SUMIFS" in moms_formula
    excel_rev = _eval_sumifs(ws, rev_formula)
    excel_moms = _eval_sumifs(ws, moms_formula)
    excel_gift = _eval_sumifs(ws, ws.cell(row=totals_row, column=10).value)

    # Oversigt — the same numbers.
    s1 = wb["Oversigt"]
    kv = {s1.cell(row=r, column=1).value: s1.cell(row=r, column=2).value
          for r in range(1, s1.max_row + 1)}
    assert kv["Omsætning i alt"] == excel_rev == 3500.0
    assert kv["Salgsmoms i alt"] == excel_moms == 700.0
    assert kv["Låste lukninger (medregnet)"] == 2
    assert kv["Kladder (ikke medregnet)"] == 1
    assert excel_gift == 500.0 and kv["Gavekort"] == 500.0

    # Range PDF — the same numbers.
    txt = pdf_text(build_daily_close_range_pdf(
        closes, from_date=f, to_date=t, business_name="Mirabelle ApS", currency="DKK"))
    assert "3.500,00 kr." in txt and "700,00 kr." in txt
    assert "30.053,00 kr." in txt   # the draft row is listed…
    assert "33.553" not in txt       # …but never summed with the locked days

    # E-mail body — the same numbers, and the draft is named.
    html = _accountant_email_body(
        business_name="Mirabelle ApS", from_date=f, to_date=t,
        totals=period_totals(closes), currency="DKK", fmt="xlsx",
        message=None, is_danish=True, attachment_name="Kasserapporter x.xlsx")
    assert "3.500,00 kr." in html and "700,00 kr." in html
    assert "1 kladde i perioden er ikke låst" in html
    assert "33.553" not in html


def test_xlsx_money_format_is_invariant_and_columns_are_wide_enough():
    """'#.##0,00' made Excel show '15021,000 kr.' and '####'. The invariant
    code is localised by Excel to '15.021,00 kr.' on a Danish machine."""
    from openpyxl import load_workbook
    raw = build_daily_close_range_xlsx(
        [_close(date(2026, 9, 1), 1234567.89, 246913.58)],
        from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        business_name="Cafe", currency="DKK")
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Kasserapport"]
    assert ws.cell(row=2, column=4).number_format == '#,##0.00" kr."'
    assert ws.cell(row=3, column=4).number_format == '#,##0.00" kr."'
    assert ws.column_dimensions["D"].width >= 16  # "1.234.567,89 kr."
    import zipfile
    styles = zipfile.ZipFile(io.BytesIO(raw)).read("xl/styles.xml").decode()
    assert "#.##0,00" not in styles


# ─── 2. The CSV opens correctly in Danish Excel ───────────────────────


def test_default_csv_parses_with_semicolon_and_decimal_comma():
    raw = closes_to_csv_bytes(_sep_period(), currency="DKK")
    assert raw.startswith("﻿".encode("utf-8"))
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    header = rows[0]
    # Danish headers, no raw DB field names, no internal ids / storage keys.
    # Each row starts with the day's own kasserapport bilag number and id.
    assert header[:7] == ["Dato", "Bilagsnr.", "Dokument-id", "Status", "Bogføring",
                          "Afdeling", "Omsætning inkl. moms"]
    for raw_name in ("revenue_total", "branch_id", "receipt_photo", "moms_total"):
        assert raw_name not in header
    money_cols = [i for i, h in enumerate(header) if h in (
        "Omsætning inkl. moms", "Salgsmoms", "Omsætning ekskl. moms", "Kontant",
        "Kort", "Gavekort", "Betalinger i alt", "Kassedifference")]
    for row in rows[1:]:
        for i in money_cols:
            cell = row[i]
            assert cell == "" or re.fullmatch(r"-?\d+,\d{2}", cell), (header[i], cell)
            assert not re.fullmatch(r"-?\d+\.\d{2}", cell)
    col = {h: i for i, h in enumerate(header)}
    body = {r[0]: r for r in rows[1:-1]}
    assert body["2026-09-26"][col["Status"]] == "Kladde (ikke medregnet)"
    assert body["2026-09-25"][col["Bilagsnr."]] == "KR-20260925-20260925"
    # Totals row: locked closes only, parses back to the period total.
    total = rows[-1]
    assert total[col["Status"]].startswith("I alt — 2 låste dage")
    assert "1 kladde ikke medregnet" in total[col["Status"]]
    assert float(total[col["Omsætning inkl. moms"]].replace(",", ".")) == 3500.0
    assert float(total[col["Salgsmoms"]].replace(",", ".")) == 700.0
    # Same figure as every other artifact.
    assert period_totals(_sep_period())["revenue"] == 3500.0


def test_machine_csv_is_only_behind_an_explicit_variant():
    default = closes_to_csv_bytes([_close(date(2026, 9, 1), 15021.0, 3004.2)]).decode("utf-8-sig")
    machine = closes_to_csv_bytes([_close(date(2026, 9, 1), 15021.0, 3004.2)],
                                  variant="machine").decode("utf-8-sig")
    assert "15021,00" in default and "15021.00" not in default
    assert machine.splitlines()[0].startswith("date;status;revenue_total")
    assert "15021.00" in machine


# ─── 4. Escaping user text in the revisor mail ────────────────────────


class _Row:
    def __init__(self, **kw):
        self.id = uuid.uuid4()
        self.date = date(2026, 9, 25)
        self.status = "confirmed"
        self.closed_at = datetime(2026, 9, 25, 23, 28)
        self.revenue_categories = "food:10000"
        self.revenue_total = 10000.0
        self.payment_categories = "cash:10000"
        self.payment_total = 10000.0
        self.moms_total = 2000.0
        self.revenue_ex_moms = 8000.0
        self.moms_mode = "auto"
        self.cash_counted = 10000.0
        self.cash_expected = 10000.0
        self.cash_difference = 0.0
        self.tips_total = None
        for k, v in kw.items():
            setattr(self, k, v)


def test_lock_mail_escapes_a_b_tag_in_lukket_af_and_the_business_name():
    from app.routers.daily_close import _build_close_email_html
    subject, html = _build_close_email_html(
        business_name='Café <img src="https://x.invalid/p.gif">', dc=_Row(),
        currency="DKK", closed_by="<b>Lars</b><a href='https://x.invalid'>bilag</a>",
        has_scan=False, scan_degraded=False, is_danish=True,
    )
    assert "&lt;b&gt;Lars&lt;/b&gt;" in html
    assert "<b>Lars" not in html
    assert "<a href='https://x.invalid'" not in html
    assert "<img" not in html
    assert "\n" not in subject and "\r" not in subject


def test_period_mail_escapes_business_name_and_message():
    from app.routers.daily_close import _accountant_email_body
    html = _accountant_email_body(
        business_name="<b>Evil</b>", from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        totals=period_totals(_sep_period()), currency="DKK", fmt="pdf",
        message="<script>x</script>", is_danish=True)
    assert "<b>Evil</b>" not in html and "&lt;b&gt;Evil&lt;/b&gt;" in html
    assert "<script>" not in html


# ─── 5. send-to-accountant mails the saved revisor only ───────────────


@pytest.fixture
def db_session(monkeypatch):
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    s = SessionLocal()

    def _override_get_db():
        try:
            yield s
        finally:
            pass

    app.dependency_overrides[get_db] = _override_get_db
    monkeypatch.setattr("app.services.billing.SessionLocal", SessionLocal, raising=False)
    import app.database as _db_mod
    monkeypatch.setattr(_db_mod, "SessionLocal", SessionLocal, raising=False)
    from app.routers import daily_close as _dc
    _dc._limiter.reset()
    try:
        yield s
    finally:
        s.close()
        app.dependency_overrides.pop(get_db, None)
        _dc._limiter.reset()


@pytest.fixture
def client():
    yield TestClient(app)
    app.dependency_overrides.clear()


def _make_user(db, *, plan="starter", email="anders@mirabelle.dk"):
    u = User(
        email=email, password_hash=hash_password("x"),
        business_name="Mirabelle Café", business_type="restaurant",
        currency="DKK", plan=plan, auto_email_on_close=True,
        email_verified=True, created_at=utc_now() - timedelta(days=2),
    )
    db.add(u); db.commit(); db.refresh(u)
    return u


def _make_profile(db, user, **kw):
    p = BusinessProfile(
        user_id=user.id, company_name="Mirabelle ApS", org_number="39842851",
        country="DK", email="owner@mirabelle.dk",
        accountant_email=kw.pop("accountant_email", "anna@revisor.dk"),
        accountant_name=kw.pop("accountant_name", "Anna"),
        **kw,
    )
    db.add(p); db.commit(); db.refresh(p)
    return p


def _auth(user):
    return {"Authorization": f"Bearer {create_access_token(str(user.id))}"}


def _add_close(db, user, d, rev=1000.0, status="confirmed"):
    c = _close(d, rev, round(rev / 5, 2), status=status)
    c.user_id = user.id
    db.add(c); db.commit(); db.refresh(c)
    return c


def test_send_to_accountant_refuses_any_address_but_the_saved_revisor(db_session, client):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    _add_close(db_session, user, date(2026, 9, 25))
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        r = client.post(
            "/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
            json={"fmt": "pdf", "accountant_email": "victim@elsewhere.dk",
                  "message": "hi"},
            headers=_auth(user))
        assert r.status_code == 422, r.text
        assert r.json()["detail"]["code"] == "recipient_not_saved"
        assert sender.call_count == 0

        r = client.post(
            "/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
            json={"fmt": "xlsx", "message": "<b>se venligst</b>"}, headers=_auth(user))
        assert r.status_code == 200, r.text
        # The revisor's message (with their opt-out) and the owner's copy are
        # SEPARATE: an owner pressing Unsubscribe in their own inbox must not
        # switch off their revisor.
        assert sender.call_count == 2
        (args, kwargs), (oargs, okwargs) = sender.call_args_list
        assert args[0] == "anna@revisor.dk"
        assert "List-Unsubscribe" in kwargs["headers"]
        # The owner's copy goes to the SAME owner address the revisor's
        # Reply-To names: the LOGIN, never the unverified Profile e-mail.
        assert oargs[0] == "anders@mirabelle.dk"
        assert kwargs["reply_to"] == "anders@mirabelle.dk"
        assert not okwargs.get("headers")
        assert "afmelde" not in oargs[2] and "Din kopi" in oargs[2]
        assert "<b>se venligst</b>" not in args[2]
        assert "&lt;b&gt;se venligst&lt;/b&gt;" in args[2]
        # Says what is attached, under the business's own name.
        assert kwargs["attachment_filename"] == "Kasserapporter Mirabelle ApS 2026-09-01–2026-09-30.xlsx"
        assert "Kasserapporter Mirabelle ApS 2026-09-01–2026-09-30.xlsx" in args[2]


def test_send_to_accountant_has_a_daily_cap_per_account(db_session, client):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    _add_close(db_session, user, date(2026, 9, 25))
    for _ in range(20):
        db_session.add(AuditLog(
            user_id=user.id, action="daily_close.send_to_accountant",
            entity_type="daily_close_range", created_at=utc_now() - timedelta(hours=1)))
    db_session.commit()
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        r = client.post(
            "/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
            json={"fmt": "pdf"}, headers=_auth(user))
    assert r.status_code == 429
    assert r.json()["detail"]["code"] == "revisor_daily_cap"
    assert sender.call_count == 0


# ═══════════════════════════════════════════════════════════════════════
# Stage 2 — the send flow
# ═══════════════════════════════════════════════════════════════════════


@pytest.fixture
def mailbox(monkeypatch):
    """Stub Resend. `mailbox.sent` collects payloads; set `mailbox.fail` to make
    the next sends raise. Nothing ever leaves the process."""
    class _Box:
        sent: list = []
        fail = False

    box = _Box()
    box.sent = []

    def _send(payload):
        if box.fail:
            raise RuntimeError("resend down")
        box.sent.append(payload)
        return {"id": "stub"}

    monkeypatch.setattr("app.services.email_service.resend.Emails.send", _send)
    monkeypatch.setattr("app.services.email_service.resend.api_key", "test_key")
    return box


@pytest.fixture
def unsub_db(db_session, monkeypatch):
    """The public unsubscribe router opens its own SessionLocal."""
    from sqlalchemy.orm import sessionmaker
    SL = sessionmaker(bind=db_session.get_bind(), autoflush=False, autocommit=False)
    monkeypatch.setattr("app.routers.email_unsubscribe.SessionLocal", SL)
    return db_session


def _lock(client, user, d="2026-09-25", rev=12500.0, **kw):
    payload = {
        "date": d, "branch_id": None, "status": "confirmed",
        "revenue_breakdown": {"food": rev},
        "payment_breakdown": {"cash": 2500.0, "card": rev - 2500.0},
        "moms_mode": "auto", "cash_counted": 2525.0,
        "closed_by": "Lars", "acknowledge_anomaly": True,
    }
    payload.update(kw)
    return client.post("/api/daily-close", json=payload, headers=_auth(user))


def _revisor_mails(box, addr="anna@revisor.dk"):
    return [p for p in box.sent if p["to"] == [addr]]


def test_failed_lock_mail_is_persisted_and_resend_is_idempotent(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    mailbox.fail = True
    r = _lock(client, user)
    assert r.status_code == 200, r.text
    close = r.json()
    assert close["close_ritual"]["email_status"] == "send_failed"
    cid = close["id"]

    # Persisted: a reload of History still knows it was NOT sent.
    listed = client.get("/api/daily-close", headers=_auth(user)).json()
    row = next(x for x in listed if x["id"] == cid)
    assert row["email_status"] == "send_failed"
    assert row["email_sent_at"] is None

    # The old retry (re-POST the locked close) is the 409 dead end; the real
    # resend works on the locked close.
    mailbox.fail = False
    r1 = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-0001"}, headers=_auth(user))
    assert r1.status_code == 200, r1.text
    assert r1.json()["replayed"] is False
    assert r1.json()["email_status"] == "sent"
    assert "anna@revisor.dk" in r1.json()["email_sent_to"]
    n_after_first = len(mailbox.sent)
    assert len(_revisor_mails(mailbox)) == 1

    # Same click again (double tap / network retry): no second mail.
    r2 = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-0001"}, headers=_auth(user))
    assert r2.status_code == 200
    assert r2.json()["replayed"] is True
    assert len(mailbox.sent) == n_after_first

    # A NEW click on an already-sent close must be confirmed first.
    r3 = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-0002"}, headers=_auth(user))
    assert r3.status_code == 409
    assert r3.json()["detail"]["code"] == "already_sent"
    assert len(mailbox.sent) == n_after_first
    r4 = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-0003", "force": True}, headers=_auth(user))
    assert r4.status_code == 200
    assert len(_revisor_mails(mailbox)) == 2


def test_resend_refuses_a_draft(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, status="draft")
    cid = r.json()["id"]
    r = client.post(f"/api/daily-close/{cid}/resend-email",
                    json={"key": "click-0001"}, headers=_auth(user))
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "not_locked"
    assert mailbox.sent == []


def test_revisor_opt_out_stops_the_next_send_and_the_owner_sees_it(
        db_session, unsub_db, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, d="2026-09-24")
    assert r.json()["close_ritual"]["accountant_included"] is True
    rev = _revisor_mails(mailbox)[0]
    # Why they get it, and the one-click opt-out — in the body and the header.
    assert "har angivet dig som revisor i BonBox" in rev["html"]
    m = re.search(r"<(https?://[^>]+/api/email/unsubscribe\?token=[^>]+)>",
                  rev["headers"]["List-Unsubscribe"])
    assert m and rev["headers"]["List-Unsubscribe-Post"] == "List-Unsubscribe=One-Click"
    token = m.group(1).split("token=", 1)[1]
    assert token in rev["html"]

    # GET only confirms (link scanners); POST opts out.
    g = client.get(f"/api/email/unsubscribe?token={token}")
    assert g.status_code == 200 and "Afmeld" in g.text
    db_session.expire_all()
    assert db_session.query(BusinessProfile).first().accountant_opted_out_at is None
    p = client.post(f"/api/email/unsubscribe?token={token}")
    assert p.status_code == 200 and "afmeldt" in p.text
    db_session.expire_all()
    prof = client.get("/api/business", headers=_auth(user)).json()
    assert prof["accountant_opted_out"] is True
    assert prof["accountant_auto_send_effective"] is False

    # The next lock mails the owner only, and says why.
    before = len(_revisor_mails(mailbox))
    r = _lock(client, user, d="2026-09-25")
    ritual = r.json()["close_ritual"]
    assert ritual["accountant_skip_reason"] == "opted_out"
    assert ritual["accountant_included"] is False
    assert len(_revisor_mails(mailbox)) == before
    owner_copy = [x for x in mailbox.sent if x["to"] == ["anders@mirabelle.dk"]][-1]
    assert "har afmeldt mails fra BonBox" in owner_copy["html"]
    # Manual sends refuse too.
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        s = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "pdf"}, headers=_auth(user))
    assert s.status_code == 409 and s.json()["detail"]["code"] == "accountant_opted_out"
    assert sender.call_count == 0
    cid = r.json()["id"]
    rr = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-0009"}, headers=_auth(user))
    assert rr.status_code == 409 and rr.json()["detail"]["code"] == "accountant_opted_out"


def test_relock_after_unlock_sends_a_marked_correction(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, rev=12500.0)
    cid = r.json()["id"]
    first = _revisor_mails(mailbox)[-1]
    assert first["subject"] == "Kasserapport fre. 25.09.2026 — Mirabelle ApS"

    u = client.post(f"/api/daily-close/{cid}/unlock",
                    json={"reason": "Forkert kortbeløb"}, headers=_auth(user))
    assert u.status_code == 200
    r = _lock(client, user, rev=13000.0)
    assert r.status_code == 200, r.text
    assert r.json()["close_ritual"]["correction"] is True
    second = _revisor_mails(mailbox)[-1]
    assert second["subject"] == "Rettet kasserapport fre. 25.09.2026 — Mirabelle ApS"
    assert "erstatter den, der blev sendt" in second["html"]
    assert "Forkert kortbeløb" in second["html"]
    assert "Omsætning 12.500,00 kr. → 13.000,00 kr." in _changes(second["html"])


def test_send_status_is_persisted_on_the_close(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user)
    cid = r.json()["id"]
    got = client.get(f"/api/daily-close/{cid}", headers=_auth(user)).json()
    assert got["email_status"] == "sent"
    assert got["email_sent_at"] is not None
    assert set(got["email_sent_to"]) == {"anna@revisor.dk", "anders@mirabelle.dk"}


def test_a_new_revisor_address_does_not_switch_on_auto_send(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user, accountant_email=None)
    bad = client.put("/api/business", json={"accountant_email": "anna@revisor"},
                     headers=_auth(user))
    assert bad.status_code == 422
    ok = client.put("/api/business", json={"accountant_email": "Anna@Revisor.dk"},
                    headers=_auth(user))
    assert ok.status_code == 200, ok.text
    assert ok.json()["accountant_email"] == "anna@revisor.dk"
    assert ok.json()["accountant_auto_send"] is False
    r = _lock(client, user)
    ritual = r.json()["close_ritual"]
    assert ritual["accountant_skip_reason"] == "auto_send_off"
    assert _revisor_mails(mailbox) == []
    # The owner ticks the explained choice → the next lock mails the revisor.
    client.put("/api/business", json={"accountant_auto_send": True}, headers=_auth(user))
    r = _lock(client, user, d="2026-09-26")
    assert r.json()["close_ritual"]["accountant_included"] is True
    assert len(_revisor_mails(mailbox)) == 1


def test_a_failed_pdf_build_is_named_not_blamed_on_the_environment(
        db_session, client, mailbox, monkeypatch):
    user = _make_user(db_session)
    _make_profile(db_session, user)

    def _boom(*a, **kw):
        raise RuntimeError("reportlab exploded")
    monkeypatch.setattr("app.routers.daily_close._close_attachment", _boom)
    r = _lock(client, user)
    ritual = r.json()["close_ritual"]
    assert ritual["email_status"] == "failed_skipped"
    assert ritual["email_error"] == "pdf_build_failed"
    assert r.json()["email_error"] == "pdf_build_failed"
    assert mailbox.sent == []


def test_migration_and_sqlite_mirror_carry_every_new_column():
    """create_all() hides a missing mirror on a fresh sqlite db, so pin both
    the Postgres ALTER and the SQLite mirror for every Migration 082/083 column."""
    import inspect
    import app.main as m
    src = inspect.getsource(m)
    for table, col in (
        ("daily_closes", "email_status"), ("daily_closes", "email_error"),
        ("daily_closes", "email_attempt_at"), ("daily_closes", "email_sent_at"),
        ("daily_closes", "email_sent_to"), ("daily_closes", "email_send_key"),
        ("business_profiles", "accountant_auto_send"),
        ("business_profiles", "accountant_opted_out_at"),
        ("business_profiles", "accountant_opted_out_email"),
        ("daily_closes", "cash_float"), ("daily_closes", "source_meta"),
    ):
        assert f"ALTER TABLE {table} ADD COLUMN IF NOT EXISTS {col} " in src, (table, col)
        assert f'_add("{table}", "{col}"' in src, ("mirror", table, col)


# ═══════════════════════════════════════════════════════════════════════
# Stage 3 — document quality (every assertion reads the rendered file)
# ═══════════════════════════════════════════════════════════════════════


def _sep25(db, user, **kw):
    """The rating's 25 Sep close: 3 categories, 4 methods, +25 drawer with a
    1.000 kr. float, tips, a note — the one that spilled onto an orphan page."""
    base = dict(
        id=uuid.uuid4(), user_id=user.id, branch_id=None, date=date(2026, 9, 25),
        revenue_categories="food:17393|drinks:9058|takeaway:2018",
        revenue_total=28469.0,
        payment_categories="cash:7971|card:16512|mobilepay:3416|gift_card:570",
        payment_total=28469.0, moms_total=5693.8, revenue_ex_moms=22775.2,
        moms_mode="auto", cash_expected=7971.0, cash_counted=7996.0,
        cash_difference=25.0, cash_float=1000.0, tips_total=832.0,
        tips_staff_count=4, tips_per_person=208.0, status="confirmed",
        notes="Travl fredag. Kortterminal 2 genstartet kl. 20.",
        closed_by="Lars", closed_at=datetime(2026, 9, 25, 23, 28), is_deleted=False,
    )
    base.update(kw)
    c = DailyClose(**base)
    db.add(c); db.commit(); db.refresh(c)
    return c


def test_single_kasserapport_carries_identity_traceability_and_fits_one_page(db_session):
    from pypdf import PdfReader
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    user = _make_user(db_session)
    prof = _make_profile(db_session, user, address="Vestergade 1", zipcode="1456", city="København K")
    dc = _sep25(db_session, user)
    out = build_close_kasserapport_pdf(db_session, user, dc, profile=prof)
    reader = PdfReader(io.BytesIO(out["pdf"]))
    assert len(reader.pages) == 1, "a typical close fits one A4 page"
    txt = pdf_text(out["pdf"])
    # Identity + traceability
    assert "Mirabelle ApS" in txt and "CVR 39842851" in txt
    assert "Fredag 25. september 2026" in txt
    # The issued number never changes (a revisor files it as the voucher).
    assert "Bilagsnr. KR-20260925-20260925" in txt
    assert f"Dokument-id: {out['doc_id']}" in txt
    assert "Side 1 af 1" in txt
    assert "Opbevares i 5 år efter bogføringsloven." in txt
    # Lock time in Copenhagen time, same as the app and the mail (not 23:28 UTC).
    assert "Låst 26.09.2026 kl. 01:28" in txt
    assert "23:28" not in txt
    # Cash as the app counts it — float shown, words for the difference.
    assert "Forventet kontant (kontantsalg)" in txt and "fra bilag" not in txt
    assert "Byttepenge" in txt and "1.000,00 kr." in txt
    assert "Optalt i skuffen i alt" in txt and "8.996,00 kr." in txt
    assert "Optalt (uden byttepenge)" in txt and "7.996,00 kr." in txt
    assert "+25,00 kr." in txt
    assert "Der er 25,00 kr. for meget i kassen." in txt
    # Not "afstemt" beside a non-zero difference.
    assert "og afstemt" not in txt
    assert "inden for tolerancen på ±100 kr." in txt
    # File name: business + date.
    assert out["filename"] == "Kasserapport Mirabelle ApS 2026-09-25.pdf"
    # Stable id: the same close renders the same document id.
    assert build_close_kasserapport_pdf(db_session, user, dc, profile=prof)["doc_id"] == out["doc_id"]


def test_single_kasserapport_names_branch_source_tills_and_history(db_session, client):
    import json as _json
    from app.models.branch import Branch
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    br = Branch(id=uuid.uuid4(), user_id=user.id, name="Mirabelle Vesterbro")
    db_session.add(br); db_session.commit()
    dc = _sep25(db_session, user, branch_id=br.id, source_meta=_json.dumps({
        "kind": "zbon", "scans": 2, "terminal_totals": [12000.0, 16469.0],
        "corrected": ["pay:card"],
    }))
    # Unlock + relock, through the audit trail the router writes.
    for action, after, when in (
        ("daily_close.lock", {"status": "confirmed"}, datetime(2026, 9, 25, 23, 28)),
        ("daily_close.unlock", {"unlock_reason": "Forkert kortbeløb", "unlocked_by": "ejer@mirabelle.dk"},
         datetime(2026, 9, 29, 7, 0)),
        ("daily_close.lock", {"status": "confirmed"}, datetime(2026, 9, 29, 7, 12)),
    ):
        db_session.add(AuditLog(user_id=user.id, action=action, entity_type="daily_close",
                                entity_id=dc.id, after_state=_json.dumps(after), created_at=when))
    db_session.commit()
    txt = pdf_text(build_close_kasserapport_pdf(db_session, user, dc, profile=prof)["pdf"])
    flat = " ".join(txt.split())
    assert "Afdeling: Mirabelle Vesterbro" in flat
    assert "Z-bon (scannet)" in flat
    assert "2 terminaler lagt sammen: 12.000,00 kr. + 16.469,00 kr." in flat
    assert "rettet af ejeren efter scanning: Kort" in flat
    assert "HISTORIK" in flat
    # The owner's role, never the login e-mail, on a document a revisor gets.
    assert "Låst op 29.09.2026 kl. 09:00 af ejeren — årsag: Forkert kortbeløb" in flat
    assert "ejer@mirabelle.dk" not in flat
    assert "Låst igen 29.09.2026 kl. 09:12" in flat

    # …and the same history and source in the Excel and the CSV.
    from openpyxl import load_workbook
    from app.routers.daily_close import _range_extras
    extras = _range_extras(db_session, user, [dc])
    wb = load_workbook(io.BytesIO(build_daily_close_range_xlsx(
        [dc], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        business_name="Mirabelle ApS", currency="DKK",
        tz=extras["tz"], history=extras["history"], sources=extras["sources"])))
    ws = wb["Kasserapport"]
    hdr = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    row = dict(zip(hdr, [ws.cell(row=2, column=i).value for i in range(1, ws.max_column + 1)]))
    assert "Forkert kortbeløb" in row["Historik"] and "Låst igen" in row["Historik"]
    assert "2 terminaler lagt sammen" in row["Kilde"]
    assert row["Låst (dansk tid)"] == datetime(2026, 9, 26, 1, 28)
    assert row["Gavekort"] == 570
    assert row["Kategori: Mad"] == 17393 and row["Kategori: Drikkevarer"] == 9058
    csv_txt = closes_to_csv_bytes([dc], currency="DKK", **extras).decode("utf-8-sig")
    assert "26.09.2026 kl. 01:28" in csv_txt
    assert "Forkert kortbeløb" in csv_txt and "Mirabelle Vesterbro" in csv_txt


def test_draft_kasserapport_says_counted_by_and_not_locked(db_session):
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    dc = _sep25(db_session, user, status="draft", closed_at=None)
    out = build_close_kasserapport_pdf(db_session, user, dc, profile=prof)
    txt = pdf_text(out["pdf"])
    assert "KLADDE" in txt
    assert "Optalt af: Lars" in txt and "Lukket af" not in txt
    assert "Ikke låst" in txt
    assert out["filename"] == "Kasserapport KLADDE Mirabelle ApS 2026-09-25.pdf"


def test_lock_mail_attaches_the_same_kasserapport_as_history(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, d="2026-09-24")
    cid = r.json()["id"]
    mailed = _revisor_mails(mailbox)[-1]["attachments"][0]
    import base64
    mailed_txt = pdf_text(base64.b64decode(mailed["content"]))
    hist = client.get(f"/api/daily-close/{cid}/pdf", headers=_auth(user))
    assert hist.status_code == 200
    assert "Kasserapport%20Mirabelle%20ApS%202026-09-24.pdf" in hist.headers["content-disposition"]
    hist_txt = pdf_text(hist.content)
    doc_id = hist.headers["x-document-id"]
    # Same document: same id, same verdict, same figures — and the mail names it.
    assert f"Dokument-id: {doc_id}" in mailed_txt and f"Dokument-id: {doc_id}" in hist_txt
    for heading in ("KLAR TIL BOGFØRING", "GENNEMGÅS"):
        assert (heading in mailed_txt) == (heading in hist_txt)
    assert mailed["filename"] == "Kasserapport Mirabelle ApS 2026-09-24.pdf"
    assert (f"Vedhæftet: Kasserapport Mirabelle ApS 2026-09-24.pdf (PDF, bilagsnr. KR-20260924-20260924, "
            f"dokument-id {doc_id})") in _revisor_mails(mailbox)[-1]["html"]


def test_range_pdf_uses_the_kasserapport_readiness_rule_and_reads_right_way_round():
    good = _close(date(2026, 9, 1), 4245.0, 849.0, pay={"cash": 4245.0},
                  cash_expected=4205.0, cash_counted=4245.0, cash_difference=40.0)
    short = _close(date(2026, 9, 2), 5000.0, 1000.0, pay={"cash": 5000.0},
                   cash_expected=5180.0, cash_counted=5000.0, cash_difference=-180.0)
    from app.services.kasserapport_claims import build_close_claims
    assert build_close_claims(good)["assurance"]["all_ok"] is True
    assert build_close_claims(short)["assurance"]["all_ok"] is False
    txt = " ".join(pdf_text(build_daily_close_range_pdf(
        [good, short], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        business_name="Cafe", currency="DKK")).split())
    assert "1 af 2 klar til bogføring · 1 skal gennemgås: 2. sep 2026" in txt
    # optalt − forventet = difference, in the app's words.
    assert "optalt 4.245,00 kr. − forventet 4.205,00 kr. = +40,00 kr. (for meget i kassen)" in txt
    assert "optalt 5.000,00 kr. − forventet 5.180,00 kr. = -180,00 kr. (kassen mangler)" in txt


def test_range_pdf_shows_gavekort_and_the_category_split():
    c = _close(date(2026, 9, 25), 28469.0, 5693.8,
               pay={"cash": 7971.0, "card": 16512.0, "mobilepay": 3416.0, "gift_card": 570.0},
               cats={"food": 17393.0, "drinks": 9058.0, "takeaway": 2018.0})
    txt = " ".join(pdf_text(build_daily_close_range_pdf(
        [c], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        business_name="Cafe", currency="DKK")).split())
    assert "Gavekort" in txt and "570,00 kr." in txt
    assert ("Omsætning pr. kategori: Mad 17.393,00 kr. · Drikkevarer 9.058,00 kr. · "
            "Takeaway 2.018,00 kr.") in txt
    assert "Betalinger pr. metode: Kontant 7.971,00 kr. · Kort 16.512,00 kr. · MobilePay 3.416,00 kr. · Gavekort 570,00 kr." in txt
    assert "1 låst · 0 kladder" in txt
    assert "Side 1 af 1" in txt
    assert "Opbevares i 5 år efter bogføringsloven" in txt and "§10" not in txt


def test_source_meta_and_float_are_saved_from_the_close_form(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, cash_float=1000.0, source_meta={
        "kind": "zbon", "scans": 2, "terminal_totals": [5000, 7500],
        "corrected": ["pay:card"], "evil": "<script>"})
    assert r.status_code == 200, r.text
    dc = db_session.query(DailyClose).filter(DailyClose.id == uuid.UUID(r.json()["id"])).first()
    assert float(dc.cash_float) == 1000.0
    import json as _json
    meta = _json.loads(dc.source_meta)
    assert meta == {"kind": "zbon", "scans": 2, "terminal_totals": [5000.0, 7500.0],
                    "corrected": ["pay:card"]}


def test_a_failed_send_after_relock_does_not_keep_saying_sent(db_session, client, mailbox):
    """History must describe THIS version: a re-lock whose mail fails cannot
    keep showing the earlier version's 'Sendt til revisor'."""
    user = _make_user(db_session)
    _make_profile(db_session, user)
    cid = _lock(client, user).json()["id"]
    client.post(f"/api/daily-close/{cid}/unlock", json={"reason": "Ret kort"}, headers=_auth(user))
    mailbox.fail = True
    r = _lock(client, user, rev=13000.0)
    assert r.json()["close_ritual"]["email_status"] == "send_failed"
    row = client.get(f"/api/daily-close/{cid}", headers=_auth(user)).json()
    assert row["email_status"] == "send_failed"
    assert row["email_sent_to"] == [] and row["email_sent_at"] is None


# ═══════════════════════════════════════════════════════════════════════
# Review fixes — correction marking from the send history, the lock mail
# inside the cap, CSRF on the opt-out, every opt-out remembered, branches,
# formula-safe Excel, no stale till list, a float that cannot 422 a lock
# ═══════════════════════════════════════════════════════════════════════


def _unlock(client, user, cid, reason="Forkert kortbeløb"):
    r = client.post(f"/api/daily-close/{cid}/unlock", json={"reason": reason}, headers=_auth(user))
    assert r.status_code == 200, r.text


def test_resend_after_a_failed_relock_is_marked_as_a_correction(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    cid = _lock(client, user, rev=12500.0).json()["id"]
    assert _revisor_mails(mailbox)[-1]["subject"] == "Kasserapport fre. 25.09.2026 — Mirabelle ApS"
    _unlock(client, user, cid)
    mailbox.fail = True
    r = _lock(client, user, rev=13000.0)
    assert r.json()["close_ritual"]["email_status"] == "send_failed"
    mailbox.fail = False
    rr = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-r001"}, headers=_auth(user))
    assert rr.status_code == 200, rr.text
    second = _revisor_mails(mailbox)[-1]
    assert second["subject"].startswith("Rettet kasserapport fre. 25.09.2026")
    assert "erstatter den, der blev sendt" in second["html"]
    assert "Forkert kortbeløb" in second["html"]
    assert "Omsætning 12.500,00 kr. → 13.000,00 kr." in _changes(second["html"])


def test_a_revisor_who_never_got_v1_is_not_told_it_is_replaced(db_session, client, mailbox):
    """auto-send off: only the owner got v1. After unlock/relock the revisor's
    FIRST mail is plain; the owner's own copy is the one marked 'Rettet'. Then a
    second correction reaches the revisor marked, after their first copy."""
    user = _make_user(db_session)
    _make_profile(db_session, user, accountant_auto_send=False)
    cid = _lock(client, user, rev=12500.0).json()["id"]
    assert _revisor_mails(mailbox) == []
    _unlock(client, user, cid)
    client.put("/api/business", json={"accountant_auto_send": True}, headers=_auth(user))
    _lock(client, user, rev=13000.0)
    rev = _revisor_mails(mailbox)
    assert len(rev) == 1
    assert rev[0]["subject"] == "Kasserapport fre. 25.09.2026 — Mirabelle ApS"
    assert "erstatter" not in rev[0]["html"]
    owner = [p for p in mailbox.sent if p["to"] == ["anders@mirabelle.dk"]][-1]
    assert "Rettet kasserapport" in owner["subject"] and "erstatter" in owner["html"]

    # Auto-send off again; the revisor now holds v2. v3 goes out by hand.
    client.put("/api/business", json={"accountant_auto_send": False}, headers=_auth(user))
    _unlock(client, user, cid, reason="Drikkepenge manglede")
    r = _lock(client, user, rev=13100.0)
    assert r.json()["close_ritual"]["accountant_skip_reason"] == "auto_send_off"
    rr = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-r002"}, headers=_auth(user))
    assert rr.status_code == 200, rr.text
    last = _revisor_mails(mailbox)[-1]
    assert last["subject"].startswith("Rettet kasserapport")
    assert "Drikkepenge manglede" in last["html"]
    assert "Omsætning 13.000,00 kr. → 13.100,00 kr." in _changes(last["html"])


def test_a_payments_only_correction_lists_the_moved_lines(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    cid = _lock(client, user, rev=12500.0,
                payment_breakdown={"cash": 2500.0, "card": 10000.0},
                revenue_breakdown={"food": 8000.0, "drinks": 4500.0}).json()["id"]
    _unlock(client, user, cid, reason="Kort var MobilePay")
    _lock(client, user, rev=12500.0,
          payment_breakdown={"cash": 2500.0, "card": 4000.0, "mobilepay": 6000.0},
          revenue_breakdown={"food": 7000.0, "drinks": 5500.0})
    html = _revisor_mails(mailbox)[-1]["html"]
    assert "Tallene er de samme" not in html
    ch = _changes(html)
    assert "Kort 10.000,00 kr. → 4.000,00 kr." in ch
    assert "MobilePay 0,00 kr. → 6.000,00 kr." in ch
    assert "Mad 8.000,00 kr. → 7.000,00 kr." in ch


def test_a_correction_without_recorded_lines_never_claims_equality():
    from app.routers.daily_close import _build_close_email_html
    dc = _close(date(2026, 9, 25), 12500.0, 2500.0)
    _s, html = _build_close_email_html(
        business_name="Mirabelle ApS", dc=dc, currency="DKK", closed_by="Lars",
        has_scan=False, scan_degraded=False, is_danish=True, audience="revisor",
        correction={"prev_sent_at": datetime(2026, 9, 25, 21, 0), "unlock_reason": "x",
                    "changes": [], "lines_known": False})
    assert "Tallene er de samme" not in html
    assert "Omsætning, salgsmoms og kassedifference er uændrede" in html


def test_the_lock_mail_counts_towards_the_daily_revisor_cap(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    for _ in range(20):
        db_session.add(AuditLog(
            user_id=user.id, action="daily_close.revisor_lock_mail",
            entity_type="daily_close", created_at=utc_now() - timedelta(hours=1)))
    db_session.commit()
    r = _lock(client, user)
    ritual = r.json()["close_ritual"]
    assert ritual["accountant_skip_reason"] == "daily_cap"
    assert ritual["accountant_included"] is False
    assert _revisor_mails(mailbox) == []
    owner = [p for p in mailbox.sent if p["to"] == ["anders@mirabelle.dk"]][-1]
    assert "loftet er nået" in owner["html"]


def test_each_lock_mail_to_the_revisor_is_counted(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    cid = _lock(client, user).json()["id"]
    for i in range(3):
        _unlock(client, user, cid, reason=f"runde {i}")
        _lock(client, user, rev=12500.0 + i)
    n = db_session.query(AuditLog).filter(
        AuditLog.user_id == user.id, AuditLog.action == "daily_close.revisor_lock_mail").count()
    assert n == 4 == len(_revisor_mails(mailbox))


def test_revisor_opt_out_works_for_a_signed_in_reader_on_the_api_host(
        db_session, unsub_db, client, mailbox):
    """The confirm page's form POSTs to api.bonbox.dk with the BonBox session
    cookie and no CSRF header — it must not 403."""
    user = _make_user(db_session)
    _make_profile(db_session, user)
    _lock(client, user)
    token = re.search(r"token=([^>\s'\"]+)",
                      _revisor_mails(mailbox)[0]["headers"]["List-Unsubscribe"]).group(1)
    signed_in = TestClient(app, base_url="https://api.bonbox.dk")
    signed_in.cookies.set("bonbox_session", "some-session")
    r = signed_in.post(f"/api/email/unsubscribe?token={token}")
    assert r.status_code == 200, r.text
    assert "afmeldt" in r.text
    db_session.expire_all()
    assert db_session.query(BusinessProfile).first().accountant_opted_out_at is not None


def test_every_opt_out_is_remembered_and_the_revisor_can_undo_it(
        db_session, unsub_db, client, mailbox):
    from app.services.revisor_mail import REVISOR_TOPIC, address_fingerprint, revisor_opted_out
    from app.utils.email_unsubscribe_token import make_unsubscribe_token
    user = _make_user(db_session)
    _make_profile(db_session, user, accountant_email="a@revisor.dk")

    def tok(addr):
        return make_unsubscribe_token(str(user.id), REVISOR_TOPIC, ttl_days=30,
                                      extra={"r": address_fingerprint(addr)})

    assert client.post(f"/api/email/unsubscribe?token={tok('a@revisor.dk')}").status_code == 200
    client.put("/api/business", json={"accountant_email": "b@revisor.dk"}, headers=_auth(user))
    assert client.post(f"/api/email/unsubscribe?token={tok('b@revisor.dk')}").status_code == 200
    client.put("/api/business", json={"accountant_email": "a@revisor.dk", "accountant_auto_send": True},
               headers=_auth(user))
    db_session.expire_all()
    prof = db_session.query(BusinessProfile).first()
    assert revisor_opted_out(prof, "a@revisor.dk") and revisor_opted_out(prof, "b@revisor.dk")
    assert "a@revisor.dk" not in (prof.accountant_opted_out_email or "")  # hashes, not addresses
    r = _lock(client, user)
    assert r.json()["close_ritual"]["accountant_skip_reason"] == "opted_out"
    assert _revisor_mails(mailbox, "a@revisor.dk") == []

    # An opt-out for an address that is no longer the saved one is still kept.
    assert client.post(f"/api/email/unsubscribe?token={tok('c@revisor.dk')}").status_code == 200
    db_session.expire_all()
    assert revisor_opted_out(db_session.query(BusinessProfile).first(), "c@revisor.dk")

    # The success page offers the revisor's own undo, and it works.
    page = client.post(f"/api/email/unsubscribe?token={tok('a@revisor.dk')}").text
    assert "undo=1" in page and "Fortryd" in page
    u = client.post(f"/api/email/unsubscribe?token={tok('a@revisor.dk')}&undo=1")
    assert u.status_code == 200 and "Du får mails igen" in u.text
    db_session.expire_all()
    prof = db_session.query(BusinessProfile).first()
    assert not revisor_opted_out(prof, "a@revisor.dk")
    assert revisor_opted_out(prof, "b@revisor.dk")
    _lock(client, user, d="2026-09-26")
    assert len(_revisor_mails(mailbox, "a@revisor.dk")) == 1


def test_two_branches_on_one_day_are_two_documents(db_session, client, mailbox):
    from openpyxl import load_workbook
    from app.models.branch import Branch
    from app.routers.daily_close import _range_extras
    user = _make_user(db_session)
    _make_profile(db_session, user)
    ves = Branch(id=uuid.uuid4(), user_id=user.id, name="Vesterbro")
    nor = Branch(id=uuid.uuid4(), user_id=user.id, name="Nørrebro")
    db_session.add_all([ves, nor]); db_session.commit()
    _lock(client, user, d="2026-09-04", branch_id=str(ves.id))
    _lock(client, user, d="2026-09-04", branch_id=str(nor.id), rev=9000.0)
    mails = _revisor_mails(mailbox)
    assert len(mails) == 2
    subjects = {m["subject"] for m in mails}
    assert subjects == {"Kasserapport fre. 04.09.2026 — Mirabelle ApS · Vesterbro",
                        "Kasserapport fre. 04.09.2026 — Mirabelle ApS · Nørrebro"}
    names = {m["attachments"][0]["filename"] for m in mails}
    assert names == {"Kasserapport Mirabelle ApS Vesterbro 2026-09-04.pdf",
                     "Kasserapport Mirabelle ApS Nørrebro 2026-09-04.pdf"}
    import base64
    texts = [" ".join(pdf_text(base64.b64decode(m["attachments"][0]["content"])).split())
             for m in mails]
    bilag = {re.search(r"Bilagsnr\. (KR-[0-9A-Z-]+)", t).group(1) for t in texts}
    assert len(bilag) == 2 and all(b.startswith("KR-20260904-20260904-") for b in bilag)
    assert {b[21:24] for b in bilag} == {"VES", "NOE"}

    closes = db_session.query(DailyClose).all()
    extras = _range_extras(db_session, user, closes)
    wb = load_workbook(io.BytesIO(build_daily_close_range_xlsx(
        closes, from_date=date(2026, 9, 4), to_date=date(2026, 9, 4),
        business_name="Mirabelle ApS", currency="DKK", **{
            k: extras[k] for k in ("tz", "history", "sources", "branch_names")})))
    ws = wb["Kasserapport"]
    hdr = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    col = hdr.index("Afdeling") + 1
    assert {ws.cell(row=2, column=col).value, ws.cell(row=3, column=col).value} == {"Vesterbro", "Nørrebro"}
    ptxt = pdf_text(build_daily_close_range_pdf(
        closes, from_date=date(2026, 9, 4), to_date=date(2026, 9, 4),
        business_name="Mirabelle ApS", currency="DKK", branch_names=extras["branch_names"]))
    assert "Vesterbro" in ptxt and "Nørrebro" in ptxt

    # A period export never shares a number with a single kasserapport.
    r = client.get("/api/daily-close/export.pdf?from=2026-09-04&to=2026-09-04", headers=_auth(user))
    assert r.status_code == 200
    assert "KRP-20260904-20260904" in pdf_text(r.content)


def test_excel_free_text_cells_are_text_never_formulas():
    from openpyxl import load_workbook
    c = _close(date(2026, 9, 25), 1000.0, 200.0,
               closed_by='=HYPERLINK("https://x.example","Lars")', notes="=1+1")
    wb = load_workbook(io.BytesIO(build_daily_close_range_xlsx(
        [c], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        business_name="=WEBSERVICE(\"https://x.example\")", currency="DKK")))
    ws = wb["Kasserapport"]
    hdr = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    lukket = ws.cell(row=2, column=hdr.index("Lukket af") + 1)
    notes = ws.cell(row=2, column=hdr.index("Bemærkninger") + 1)
    assert lukket.data_type == "s" and lukket.value.startswith("=HYPERLINK")
    assert notes.data_type == "s" and notes.value == "=1+1"
    assert wb["Oversigt"]["A1"].data_type == "s"
    # Our own totals formulas stay formulas.
    assert ws.cell(row=3, column=4).data_type == "f"


def test_an_edit_after_unlock_drops_the_stale_till_list(db_session, client, mailbox):
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    cid = _lock(client, user, rev=12500.0, source_meta={
        "kind": "zbon", "scans": 2, "terminal_totals": [5000.0, 7500.0], "corrected": ["pay:card"],
    }).json()["id"]
    _unlock(client, user, cid)
    _lock(client, user, rev=13000.0, source_meta=None)
    dc = db_session.query(DailyClose).filter(DailyClose.id == uuid.UUID(cid)).first()
    db_session.refresh(dc)
    flat = " ".join(pdf_text(build_close_kasserapport_pdf(db_session, user, dc, profile=prof)["pdf"]).split())
    assert "Z-bon (scannet)" in flat
    assert "rettet af ejeren efter oplåsning" in flat
    assert "5.000,00 kr. + 7.500,00 kr." not in flat


def test_an_out_of_range_float_never_blocks_the_lock(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, cash_float=-500)
    assert r.status_code == 200, r.text
    assert r.json().get("cash_float") in (None, 0)
    r = _lock(client, user, d="2026-09-26", cash_float=5_000_000)
    assert r.status_code == 200, r.text
    r = _lock(client, user, d="2026-09-27", cash_float=1000)
    assert r.status_code == 200
    dc = db_session.query(DailyClose).filter(DailyClose.date == date(2026, 9, 27)).first()
    assert float(dc.cash_float) == 1000.0


def test_a_send_in_flight_blocks_a_second_one(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    mailbox.fail = True
    cid = _lock(client, user).json()["id"]
    mailbox.fail = False
    dc = db_session.query(DailyClose).filter(DailyClose.id == uuid.UUID(cid)).first()
    dc.email_status = "sending"; dc.email_attempt_at = utc_now(); dc.email_send_key = "other-tab-key"
    db_session.commit()
    r = client.post(f"/api/daily-close/{cid}/resend-email", json={"key": "click-x001"}, headers=_auth(user))
    assert r.status_code == 409 and r.json()["detail"]["code"] == "in_progress"
    assert _revisor_mails(mailbox) == []
    # A claim left by a crashed worker goes stale and can be taken over.
    dc.email_attempt_at = utc_now() - timedelta(minutes=10)
    db_session.commit()
    r = client.post(f"/api/daily-close/{cid}/resend-email", json={"key": "click-x002"}, headers=_auth(user))
    assert r.status_code == 200, r.text
    assert r.json()["email_status"] == "sent"
    assert len(_revisor_mails(mailbox)) == 1


def test_a_card_only_day_without_a_cash_count_is_book_ready():
    from app.services.kasserapport_claims import build_close_claims
    c = _close(date(2026, 9, 25), 12500.0, 2500.0, pay={"card": 12500.0},
               cash_expected=None, cash_counted=None, cash_difference=None)
    a = build_close_claims(c, currency="DKK", has_bilag=False, bilagsnummer="KR-1")["assurance"]
    assert a["all_ok"] is True
    assert any("Ingen kontantsalg" in x["text"] for x in a["checks"])
    txt = pdf_text(build_daily_close_range_pdf(
        [c], from_date=date(2026, 9, 25), to_date=date(2026, 9, 25),
        business_name="Webshop", currency="DKK"))
    assert "1 af 1 klar til bogføring" in txt
    # A cash day that was not counted is still flagged.
    c2 = _close(date(2026, 9, 25), 12500.0, 2500.0, pay={"cash": 2500.0, "card": 10000.0},
                cash_counted=None, cash_difference=None)
    assert build_close_claims(c2, currency="DKK")["assurance"]["all_ok"] is False
