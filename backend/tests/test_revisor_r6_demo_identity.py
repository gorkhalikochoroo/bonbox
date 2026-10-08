"""Revisor artifacts, round 6 — never mail the revisor under the sample
company; demo seeding and clearing keep the owner's own data.

1. The demo seeder writes its sample company (Mirabelle ApS, CVR 39842851,
   Vestergade 1) into the business profile. An owner who then saved their REAL
   revisor (as the demo notice tells them to) had every real day mailed to that
   revisor as "Mirabelle ApS (CVR 39842851)" — subject, From, footer and the
   kasserapport. While the identity is still the sample's, nothing goes to a
   revisor: lock (skip), resend (owner copy only), period / MOMS / payroll /
   invite (409) — reason "demo_identity", and the owner is told
   "Din virksomhed står stadig som eksempelvirksomheden (Mirabelle ApS). Ret
   navn, CVR og adresse under Profil, før vi sender til din revisor." Once the
   owner saves their own name and CVR, the sends work again.
2. A real day's document under the sample identity says so ("Virksomheds-
   oplysningerne er eksempeldata — ret dem under Profil"); a demo DAY keeps
   its EKSEMPEL treatment; a real identity's documents are byte-identical to
   what they were (the new branch never runs for them).
3. Seeding never touches a CVR-verified identity, snapshots what the owner had,
   and keeps an owner's own revisor and bank details. "Ryd demodata" resets
   only the seeded values: an owner-entered revisor and bank details survive.
4. Polish: a period send carries a key (a retried POST never mails twice), a
   "+tag" address does not step round an opt-out, and the MOMS/payroll owner
   copies honour the opt-out like the close paths.

Every send is stubbed — nothing leaves the process. The documents are read
with pypdf / openpyxl.
"""
from __future__ import annotations

import hashlib
import io
import json
import uuid
import zipfile
from datetime import date, datetime, timedelta

import pytest

from app.models.audit_log import AuditLog
from app.models.business_profile import BusinessProfile
from app.models.daily_close import DailyClose, encode_breakdown
from app.services.revisor_mail import (
    DEMO_IDENTITY_DOC_LINE_DA,
    DEMO_IDENTITY_MESSAGE_DA,
    DEMO_IDENTITY_MESSAGE_EN,
)
from tests.test_revisor_artifacts import (  # noqa: F401 — fixtures
    _auth,
    _lock,
    _make_profile,
    _make_user,
    _sep25,
    client,
    db_session,
    mailbox,
    pdf_text,
)

PIA = "pia@realrevisor.dk"
OWN_COMPANY = {"company_name": "Testcafé ApS", "org_number": "12345678"}


@pytest.fixture(autouse=True)
def _fresh_route_limiters():
    """The MOMS, payroll, invite, Profile and demo routes carry per-IP
    limiters: none may leak into (or in from) another test file."""
    from app.routers import (
        accountants as _acc, business_profile as _bp, demo as _demo, staff as _staff,
        tax as _tax,
    )
    lims = (_tax._limiter, _staff._limiter, _acc._limiter, _bp._limiter, _demo.limiter)
    for lim in lims:
        lim.reset()
    yield
    for lim in lims:
        lim.reset()


def _flat(txt: str) -> str:
    return " ".join((txt or "").replace("\xa0", " ").split())


def _seeded_owner(db, *, plan="starter", email="owner@testcafe.dk"):
    """A real owner who tried the demo: the real seeder ran on their account."""
    from app.services.demo_seed import seed_for_user
    user = _make_user(db, plan=plan, email=email)
    user.business_name = "Testcafé (signup)"
    db.commit()
    assert seed_for_user(db, user)["ok"] is True
    return user


def _profile(db, user) -> BusinessProfile:
    p = db.query(BusinessProfile).filter_by(user_id=user.id).first()
    if p is not None:
        db.refresh(p)
    return p


def _save_revisor(client, user, *, name=None):
    """Exactly what ProfilePage.saveAccountant sends: company_name echoed."""
    return client.put("/api/business", json={
        "company_name": "Mirabelle ApS", "accountant_email": PIA,
        "accountant_name": name, "accountant_auto_send": True,
    }, headers=_auth(user))


def _row(db, user, d, rev):
    c = DailyClose(
        id=uuid.uuid4(), user_id=user.id, branch_id=None, date=d,
        revenue_categories=encode_breakdown({"food": rev}), revenue_total=rev,
        payment_categories=encode_breakdown({"card": rev}), payment_total=rev,
        moms_total=round(rev / 5, 2), revenue_ex_moms=round(rev - rev / 5, 2),
        moms_mode="auto", status="confirmed", closed_by="Lars",
        closed_at=datetime.combine(d, datetime.min.time()) + timedelta(hours=21),
        is_deleted=False, notes="rigtig dag",
    )
    db.add(c); db.commit(); db.refresh(c)
    return c


def _to(box, addr):
    return [p for p in box.sent if p["to"] == [addr]]


# ═══ 1. Nothing goes to a revisor under the sample company ═══════════════


def test_probe_l_a_real_day_is_not_mailed_to_the_real_revisor_under_the_sample_company(
        db_session, client, mailbox):
    user = _seeded_owner(db_session)
    r = _save_revisor(client, user)
    assert r.status_code == 200, r.text
    prof = r.json()
    assert prof["accountant_is_demo"] is False          # a real revisor …
    assert prof["identity_is_demo"] is True             # … under the sample company
    assert prof["accountant_auto_send_effective"] is False

    lock = _lock(client, user, d="2026-10-07")          # a REAL day
    assert lock.status_code == 200, lock.text
    ritual = lock.json()["close_ritual"]
    assert ritual["accountant_skip_reason"] == "demo_identity"
    assert ritual["accountant_included"] is False
    assert PIA not in ritual["sent_to"]
    assert _to(mailbox, PIA) == []
    # The owner is told why, in their own copy.
    owner = _to(mailbox, "owner@testcafe.dk")
    assert len(owner) == 1
    assert DEMO_IDENTITY_MESSAGE_DA in _flat(owner[0]["html"])
    # And it is on the audit row.
    a = (db_session.query(AuditLog).filter(AuditLog.action == "close.auto_emailed")
         .order_by(AuditLog.created_at.desc()).first())
    assert json.loads(a.after_state)["accountant_skip_reason"] == "demo_identity"
    # The sample company's name never replaced the signup name.
    db_session.refresh(user)
    assert user.business_name == "Testcafé (signup)"


def test_resend_sends_the_revisor_nothing_under_the_sample_company(db_session, client, mailbox):
    user = _seeded_owner(db_session)
    _save_revisor(client, user)
    cid = _lock(client, user, d="2026-10-07").json()["id"]
    mailbox.sent.clear()
    rr = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-identity-1", "force": True}, headers=_auth(user))
    assert rr.status_code == 200, rr.text
    assert rr.json()["close_ritual"]["accountant_skip_reason"] == "demo_identity"
    assert _to(mailbox, PIA) == []
    assert [p["to"] for p in mailbox.sent] == [["owner@testcafe.dk"]]


def test_period_moms_payroll_and_invite_answer_409_demo_identity(db_session, client, mailbox):
    user = _seeded_owner(db_session, plan="pro")
    _save_revisor(client, user)
    _row(db_session, user, date(2026, 10, 2), 1000.0)
    for url, body in (
        ("/api/daily-close/send-to-accountant?from=2026-10-01&to=2026-10-07", {"fmt": "pdf"}),
        ("/api/tax/filing-pdf/send-to-accountant?period_start=2026-10-01&period_end=2026-10-07",
         {"cc_self": True}),
        ("/api/staff/payroll/send-to-accountant",
         {"period_start": "2026-10-01", "period_end": "2026-10-07"}),
        ("/api/accountants/invite", {"email": PIA}),
    ):
        r = client.post(url, json=body, headers=_auth(user))
        assert r.status_code == 409, (url, r.status_code, r.text)
        d = r.json()["detail"]
        assert d["code"] == "demo_identity", url
        assert d["message_da"] == DEMO_IDENTITY_MESSAGE_DA
        assert d["message"] == DEMO_IDENTITY_MESSAGE_EN
    assert mailbox.sent == []


def test_after_the_owner_saves_their_own_name_and_cvr_the_sends_work(db_session, client, mailbox):
    user = _seeded_owner(db_session, plan="pro")
    _save_revisor(client, user)
    r = client.put("/api/business", json=OWN_COMPANY, headers=_auth(user))
    assert r.status_code == 200, r.text
    prof = r.json()
    assert prof["identity_is_demo"] is False
    assert prof["accountant_auto_send_effective"] is True
    # What only the sample company had went with it: the " · demo" stamp
    # (typed by hand ≠ verified), its street, phone, e-mail and VAT number.
    p = _profile(db_session, user)
    assert p.cvr_verified_source is None and p.cvr_verified_at is None
    assert p.address is None and p.zipcode is None and p.city is None
    assert p.vat_number is None and p.phone is None and p.email is None
    assert p.accountant_email == PIA  # the owner's revisor stays
    db_session.refresh(user)
    assert user.business_name == "Testcafé ApS"

    ritual = _lock(client, user, d="2026-10-07").json()["close_ritual"]
    assert ritual["accountant_included"] is True
    to_pia = _to(mailbox, PIA)
    assert len(to_pia) == 1
    m = to_pia[0]
    assert "Testcafé ApS" in m["subject"] and "Mirabelle" not in m["subject"]
    assert "Mirabelle" not in m["html"] and "39842851" not in m["html"]
    assert "CVR 12345678" in m["html"]

    _row(db_session, user, date(2026, 10, 2), 1000.0)
    r = client.post("/api/daily-close/send-to-accountant?from=2026-10-01&to=2026-10-07",
                    json={"fmt": "pdf", "cc_self": False}, headers=_auth(user))
    assert r.status_code == 200, r.text
    t = client.post("/api/tax/filing-pdf/send-to-accountant"
                    "?period_start=2026-10-01&period_end=2026-10-07",
                    json={"cc_self": False}, headers=_auth(user))
    # Not the identity fence any more (the period still holds the demo's
    # sample days and expenses, which the MOMS send refuses on its own).
    assert not (t.status_code == 409 and t.json()["detail"].get("code") == "demo_identity"), t.text
    inv = client.post("/api/accountants/invite", json={"email": PIA}, headers=_auth(user))
    assert inv.status_code == 201, inv.text
    assert inv.json()["owner_business_name"] == "Testcafé ApS"


def test_the_name_alone_is_not_enough_the_cvr_must_be_the_owners_too(db_session, client, mailbox):
    user = _seeded_owner(db_session)
    _save_revisor(client, user)
    r = client.put("/api/business", json={"company_name": "Testcafé ApS"}, headers=_auth(user))
    assert r.status_code == 200 and r.json()["identity_is_demo"] is True
    assert _lock(client, user, d="2026-10-07").json()["close_ritual"]["accountant_skip_reason"] \
        == "demo_identity"
    assert _to(mailbox, PIA) == []


def test_a_register_save_that_echoes_the_sample_company_stays_fenced(db_session, client, mailbox):
    """BusinessLookup echoes source "cvrapi.dk" for the profile it was opened
    with: the tag is replaced, but the sample company is still there."""
    user = _seeded_owner(db_session)
    _save_revisor(client, user)
    r = client.put("/api/business", json={
        "company_name": "Mirabelle ApS", "org_number": "39842851", "address": "Vestergade 1",
        "city": "København K", "zipcode": "1456", "source": "cvrapi.dk",
    }, headers=_auth(user))
    assert r.status_code == 200, r.text
    assert r.json()["cvr_verified_source"] == "cvrapi.dk"
    assert r.json()["identity_is_demo"] is True
    assert _lock(client, user, d="2026-10-07").json()["close_ritual"]["accountant_skip_reason"] \
        == "demo_identity"


def test_the_predicate_reads_the_sample_company_not_a_name_alone():
    from app.services.revisor_mail import is_demo_identity
    P = BusinessProfile
    assert is_demo_identity(P(company_name="Mirabelle ApS", org_number="39842851",
                              cvr_verified_source="cvrapi.dk · demo"))
    assert is_demo_identity(P(company_name="Mirabelle ApS", org_number="39842851",
                              address="Vestergade 1", zipcode="1456"))
    assert is_demo_identity(P(company_name="Mit Navn", org_number="39842851",
                              address="Vestergade 1", zipcode="1456"))
    assert is_demo_identity(P(company_name="X", org_number="39842851",
                              cvr_verified_at=datetime(2026, 9, 1)))
    # A business the owner typed themselves, with no sample address or stamp.
    assert not is_demo_identity(P(company_name="Mirabelle ApS", org_number="39842851"))
    assert not is_demo_identity(P(company_name="Café Rigtig ApS", org_number="12345678",
                                  address="Vestergade 1", zipcode="1456"))
    assert not is_demo_identity(None)


# ═══ 2. The documents ═════════════════════════════════════════════════════


def _tagged_profile(db, user, **kw):
    p = BusinessProfile(
        user_id=user.id, company_name="Mirabelle ApS", org_number="39842851",
        address="Vestergade 1", zipcode="1456", city="København K", country="DK",
        accountant_email=PIA, cvr_verified_source="cvrapi.dk · demo",
        cvr_verified_at=datetime(2026, 9, 1), **kw,
    )
    db.add(p); db.commit(); db.refresh(p)
    return p


def test_a_real_day_under_the_sample_company_says_so_on_its_kasserapport(db_session):
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    user = _make_user(db_session)
    prof = _tagged_profile(db_session, user)
    out = build_close_kasserapport_pdf(db_session, user, _sep25(db_session, user), profile=prof)
    txt = _flat(pdf_text(out["pdf"]))
    assert DEMO_IDENTITY_DOC_LINE_DA in txt
    # Not a voucher claim: the day is the owner's — it keeps its number and
    # its verdict, and is not called an example.
    assert "Bilagsnr. KR-20260925-20260925" in txt
    assert "EKSEMPEL" not in txt and "ikke til bogføring" not in txt
    assert "Side 1 af 1" in txt


def test_a_demo_day_keeps_its_eksempel_treatment_without_the_identity_line(db_session):
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    user = _make_user(db_session)
    prof = _tagged_profile(db_session, user)
    dc = _sep25(db_session, user, notes="sample · demo")
    txt = _flat(pdf_text(build_close_kasserapport_pdf(db_session, user, dc, profile=prof)["pdf"]))
    assert "KASSERAPPORT — EKSEMPEL" in txt
    assert DEMO_IDENTITY_DOC_LINE_DA not in txt


def test_the_period_pdf_and_excel_under_the_sample_company_say_so(db_session):
    from openpyxl import load_workbook
    from app.services.daily_close_range_export import (
        build_daily_close_range_pdf, build_daily_close_range_xlsx,
    )
    user = _make_user(db_session)
    prof = _tagged_profile(db_session, user)
    closes = [_sep25(db_session, user)]
    kw = dict(from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
              business_name="Mirabelle ApS", currency="DKK", profile=prof)
    assert DEMO_IDENTITY_DOC_LINE_DA in _flat(pdf_text(build_daily_close_range_pdf(closes, **kw)))
    ws = load_workbook(io.BytesIO(build_daily_close_range_xlsx(closes, **kw)))["Oversigt"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert DEMO_IDENTITY_DOC_LINE_DA in cells
    # The overview below it did not move.
    assert ws["A6"].value == "Periode"


def _frozen_documents(db, user, prof, monkeypatch):
    """Every revisor document of a real identity, rendered on a frozen clock
    with reportlab's invariant mode — so two renders compare byte for byte."""
    import reportlab.rl_config as rlc
    import app.utils.time as _t
    from app.routers.daily_close import _range_bilagsnummer, _range_extras
    from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
    from app.services.daily_close_range_export import (
        build_daily_close_range_pdf, build_daily_close_range_xlsx, closes_to_csv_bytes,
    )
    monkeypatch.setattr(rlc, "invariant", 1)
    monkeypatch.setattr(_t, "utc_now", lambda: datetime(2026, 10, 8, 10, 0))
    dc = db.query(DailyClose).filter_by(user_id=user.id).order_by(DailyClose.date).all()
    f, t = date(2026, 9, 1), date(2026, 9, 30)
    ex = _range_extras(db, user, dc)
    kw = dict(from_date=f, to_date=t, business_name=prof.company_name, currency="DKK",
              profile=prof, db=db, user_id=user.id, tz=ex["tz"], history=ex["history"],
              sources=ex["sources"], branch_names=ex["branch_names"],
              bilagsnummer=_range_bilagsnummer(f, t))
    xlsx = zipfile.ZipFile(io.BytesIO(build_daily_close_range_xlsx(dc, **kw)))
    return {
        "kasserapport": build_close_kasserapport_pdf(db, user, dc[-1], profile=prof)["pdf"],
        "period.pdf": build_daily_close_range_pdf(dc, **kw),
        # docProps/core.xml holds openpyxl's own save time — every other part.
        "period.xlsx": {n: hashlib.sha256(xlsx.read(n)).hexdigest()
                        for n in xlsx.namelist() if n != "docProps/core.xml"},
        "period.csv": closes_to_csv_bytes(dc, currency="DKK", **ex),
    }


@pytest.mark.parametrize("identity", [
    dict(company_name="Café Rigtig ApS", org_number="12345678", address="Nørregade 2",
         zipcode="1165", city="København K"),
    # The repo's own test fixture: the sample name and CVR typed in, unverified.
    dict(company_name="Mirabelle ApS", org_number="39842851"),
])
def test_real_identity_documents_are_byte_identical(db_session, monkeypatch, identity):
    """Regression guard: for a real identity the new branch never runs — the
    files are byte for byte what they are with the identity check removed.
    (Checked against the pre-change commit too: same SHA-256 for all four.)"""
    import app.services.revisor_mail as rm
    user = _make_user(db_session)
    prof = BusinessProfile(user_id=user.id, country="DK", accountant_email=PIA, **identity)
    db_session.add(prof); db_session.commit(); db_session.refresh(prof)
    _sep25(db_session, user)
    _sep25(db_session, user, id=uuid.uuid4(), date=date(2026, 9, 24), notes=None)
    assert rm.is_demo_identity(prof) is False
    now = _frozen_documents(db_session, user, prof, monkeypatch)
    monkeypatch.setattr(rm, "is_demo_identity", lambda _p: False)
    without = _frozen_documents(db_session, user, prof, monkeypatch)
    assert now == without
    assert DEMO_IDENTITY_DOC_LINE_DA not in _flat(pdf_text(now["kasserapport"]))
    assert DEMO_IDENTITY_DOC_LINE_DA not in _flat(pdf_text(now["period.pdf"]))


# ═══ 3. Seeding keeps the owner's own data ════════════════════════════════


def test_seeding_never_touches_a_cvr_verified_identity(db_session):
    from app.services.demo_seed import _seed_business_profile, seed_for_user
    user = _make_user(db_session)
    p = BusinessProfile(user_id=user.id, company_name="Real Café ApS", org_number="12345678",
                        address="Nørregade 2", accountant_email=PIA, bank_reg_number="1234",
                        cvr_verified_at=datetime(2026, 9, 1), cvr_verified_source="cvrapi.dk")
    db_session.add(p); db_session.commit()
    assert seed_for_user(db_session, user)["ok"] is False
    _seed_business_profile(db_session, user, mark_demo=True)  # a bypassed gate
    db_session.commit(); db_session.refresh(p)
    assert (p.company_name, p.org_number, p.address) == ("Real Café ApS", "12345678", "Nørregade 2")
    assert p.accountant_email == PIA and p.bank_reg_number == "1234"
    assert p.cvr_verified_source == "cvrapi.dk" and p.demo_snapshot_json is None


def test_seeding_an_empty_account_still_shows_the_whole_demo(db_session):
    user = _seeded_owner(db_session)
    p = _profile(db_session, user)
    assert p.company_name == "Mirabelle ApS" and p.org_number == "39842851"
    assert p.address == "Vestergade 1" and p.cvr_verified_source == "cvrapi.dk · demo"
    assert p.accountant_email == "revisor@mirabelle.example" and p.accountant_name == "Anna Hansen"
    assert p.accountant_auto_send is False
    assert p.identity_is_demo is True
    assert json.loads(p.demo_snapshot_json)["created"] is True


def test_seeding_over_an_unverified_profile_keeps_the_owners_revisor_and_bank(db_session):
    from app.services.demo_seed import clear_for_user, seed_for_user
    user = _make_user(db_session)
    p = BusinessProfile(user_id=user.id, company_name="Min Café", address="Jægersborggade 7",
                        accountant_email=PIA, accountant_name="Pia", accountant_auto_send=True,
                        bank_reg_number="1234", bank_account_number="0012345678",
                        day_cutoff_hour=5)
    db_session.add(p); db_session.commit()
    assert seed_for_user(db_session, user)["ok"] is True
    db_session.refresh(p)
    # The demo shows its sample company …
    assert p.company_name == "Mirabelle ApS" and p.identity_is_demo is True
    # … but the owner's revisor (never mailed under it) and bank stay.
    assert p.accountant_email == PIA and p.accountant_name == "Pia"
    assert p.bank_reg_number == "1234"
    snap = json.loads(p.demo_snapshot_json)
    assert snap["fields"]["company_name"] == "Min Café"
    assert snap["fields"]["address"] == "Jægersborggade 7"

    out = clear_for_user(db_session, user)
    db_session.refresh(p)
    # The owner's own values are back, the sample's gone.
    assert (p.company_name, p.address, p.org_number) == ("Min Café", "Jægersborggade 7", None)
    assert p.day_cutoff_hour == 5
    assert p.cvr_verified_source is None and p.cvr_verified_at is None
    assert p.accountant_email == PIA and p.accountant_auto_send is True
    assert p.bank_account_number == "0012345678"
    assert p.demo_snapshot_json is None and p.identity_is_demo is False
    assert set(out["kept"]) == {"identity", "revisor", "bank"}


# ═══ 3b. Ryd demodata keeps what the owner typed ══════════════════════════


def test_probe_k_clearing_keeps_the_owners_revisor_and_bank_details(db_session, client):
    user = _seeded_owner(db_session)
    assert _save_revisor(client, user, name="Pia Jensen").status_code == 200
    assert client.put("/api/business", json={
        "company_name": "Mirabelle ApS", "bank_reg_number": "1234",
        "bank_account_number": "0012345678",
    }, headers=_auth(user)).status_code == 200

    r = client.post("/api/demo/clear", headers=_auth(user))
    assert r.status_code == 200, r.text
    assert r.json()["deleted"]["business_profile_reset"] == 1
    assert set(r.json()["kept"]) == {"revisor", "bank"}

    prof = client.get("/api/business", headers=_auth(user)).json()
    assert prof is not None
    assert prof["accountant_email"] == PIA and prof["accountant_name"] == "Pia Jensen"
    assert prof["accountant_auto_send"] is True
    assert prof["bank_reg_number"] == "1234" and prof["bank_account_number"] == "0012345678"
    # Only the seeded values were reset: no sample company, no " · demo" stamp.
    assert prof["company_name"] == "" and prof["org_number"] is None
    assert prof["address"] is None and prof["cvr_verified_source"] is None
    assert prof["cvr_verified_at"] is None and prof["identity_is_demo"] is False
    assert prof["day_cutoff_hour"] == 6
    db_session.refresh(user)
    assert user.business_name == "Testcafé (signup)"


def test_clearing_an_untouched_demo_profile_still_leaves_a_clean_slate(db_session, client):
    user = _seeded_owner(db_session)
    r = client.post("/api/demo/clear", headers=_auth(user))
    assert r.status_code == 200 and r.json()["kept"] == []
    assert client.get("/api/business", headers=_auth(user)).json() is None


def test_clearing_after_the_owner_fixed_the_company_keeps_all_of_it(db_session, client):
    user = _seeded_owner(db_session)
    _save_revisor(client, user)
    client.put("/api/business", json={**OWN_COMPANY, "address": "Nørregade 2",
                                      "industry": "Restauranter"}, headers=_auth(user))
    client.post("/api/demo/clear", headers=_auth(user))
    prof = client.get("/api/business", headers=_auth(user)).json()
    assert prof["company_name"] == "Testcafé ApS" and prof["org_number"] == "12345678"
    assert prof["address"] == "Nørregade 2" and prof["industry"] == "Restauranter"
    assert prof["accountant_email"] == PIA


def test_a_legacy_seeded_profile_without_snapshot_keeps_the_owners_revisor(db_session):
    """Seeded before snapshots existed: still only the seeded values go."""
    from app.services.demo_seed import clear_for_user
    user = _make_user(db_session)
    p = _tagged_profile(db_session, user, bank_reg_number="4321",
                        phone="+45 33 11 22 33", day_cutoff_hour=4, email="info@mirabelle.dk")
    out = clear_for_user(db_session, user)
    db_session.refresh(p)
    assert p.accountant_email == PIA and p.bank_reg_number == "4321"
    assert p.company_name == "" and p.org_number is None and p.phone is None
    assert p.email is None  # the old seed's business address goes too
    assert p.cvr_verified_source is None and p.day_cutoff_hour == 6
    assert set(out["kept"]) == {"revisor", "bank"}


# ═══ 4. Polish that reaches the revisor or the owner ══════════════════════


def test_a_retried_period_send_with_the_same_key_mails_once(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user, accountant_email=PIA, accountant_name="Pia")
    _row(db_session, user, date(2026, 9, 25), 1000.0)
    url = "/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30"
    body = {"fmt": "csv", "cc_self": False, "key": "tap-period-001"}
    a = client.post(url, json=body, headers=_auth(user))
    b = client.post(url, json=body, headers=_auth(user))
    assert a.status_code == 200 and b.status_code == 200, (a.text, b.text)
    assert b.json()["replayed"] is True and b.json()["sent_to"] == PIA
    assert b.json()["filename"] == a.json()["filename"]
    assert len(_to(mailbox, PIA)) == 1
    # A new tap (new key) is a new send.
    c = client.post(url, json={**body, "key": "tap-period-002"}, headers=_auth(user))
    assert c.status_code == 200 and "replayed" not in c.json()
    assert len(_to(mailbox, PIA)) == 2


def test_a_plus_tag_does_not_step_round_an_opt_out():
    from app.services.revisor_mail import (
        address_fingerprint, mailbox_address, owner_copy_allowed, revisor_opted_out,
    )
    p = BusinessProfile(accountant_email="pia+bonbox@realrevisor.dk",
                        accountant_opted_out_email=address_fingerprint(PIA))
    assert mailbox_address("Pia+BonBox@RealRevisor.dk") == PIA
    assert revisor_opted_out(p) is True
    assert owner_copy_allowed(p, "pia+x@realrevisor.dk") is False
    # Another mailbox at the same firm is not opted out.
    assert revisor_opted_out(p, "kim@realrevisor.dk") is False


def test_the_moms_owner_copy_honours_the_opt_out_like_the_close_paths(db_session, client):
    from unittest.mock import patch
    from app.services.revisor_mail import address_fingerprint
    user = _make_user(db_session, email="owner@testcafe.dk")
    _make_profile(db_session, user, accountant_email="min@revisor.dk",
                  accountant_opted_out_email=address_fingerprint("owner@testcafe.dk"))
    _row(db_session, user, date(2026, 9, 25), 1000.0)
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        t = client.post("/api/tax/filing-pdf/send-to-accountant"
                        "?period_start=2026-09-01&period_end=2026-09-30",
                        json={"cc_self": True}, headers=_auth(user))
    assert t.status_code == 200, t.text
    assert [c.args[0] for c in sender.call_args_list] == ["min@revisor.dk"]


def test_the_expired_opt_out_page_names_every_bonbox_mail():
    from app.routers.email_unsubscribe import _revisor_expired_page
    html = _revisor_expired_page("Testcafé ApS")
    assert "Afmeldingslinks i mails fra BonBox" in html
    assert "mails med kasserapporter" not in html


def test_migration_and_sqlite_mirror_carry_the_snapshot_column():
    import inspect
    import app.main as m
    src = inspect.getsource(m)
    assert ("ALTER TABLE business_profiles ADD COLUMN IF NOT EXISTS demo_snapshot_json TEXT"
            in src)
    assert '_add("business_profiles", "demo_snapshot_json", "TEXT")' in src
    # Never writable or readable through the API.
    from app.schemas.business_profile import BusinessProfileCreate, BusinessProfileResponse
    assert "demo_snapshot_json" not in BusinessProfileCreate.model_fields
    assert "demo_snapshot_json" not in BusinessProfileResponse.model_fields
