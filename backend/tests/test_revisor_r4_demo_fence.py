"""Revisor artifacts, round 4 — demo data never mails a third party, demo days
stay out of the revisor's period.

The demo seeder used to write a deliverable revisor (anna@revisor.dk) with
auto-send unset (read as ON): the first REAL day an owner locked after trying
the demo was mailed to an address the owner never typed. And the period send
counted the seeded " · demo" days as kasserapporter. Every send is stubbed —
nothing leaves the process.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
import uuid
from datetime import date, datetime, timedelta
from unittest.mock import patch

import pytest

from app.models.audit_log import AuditLog
from app.models.business_profile import BusinessProfile
from app.models.daily_close import DailyClose, encode_breakdown
from tests.test_revisor_artifacts import (  # noqa: F401 — fixtures
    _auth,
    _lock,
    _make_user,
    _revisor_mails,
    _unlock,
    client,
    db_session,
    mailbox,
    pdf_text,
    unsub_db,
)

DEMO_LINE_DA = "Revisoren er eksempeldata — gem din egen revisors mail under Profil"


@pytest.fixture(autouse=True)
def _fresh_route_limiters():
    """These tests call the MOMS, payroll and invite routes too: their per-IP
    minute limiters must not leak into (or in from) other test files."""
    from app.routers import accountants as _acc, staff as _staff, tax as _tax
    for lim in (_tax._limiter, _staff._limiter, _acc._limiter):
        lim.reset()
    yield
    for lim in (_tax._limiter, _staff._limiter, _acc._limiter):
        lim.reset()


def _demo_profile(db, user, *, accountant_email="anna@revisor.dk", email="info@mirabelle.dk",
                  auto_send=None):
    """A profile as the OLD demo seeder left it: tagged " · demo", the seeded
    revisor at a real domain, auto-send NULL (read as on)."""
    p = BusinessProfile(
        user_id=user.id, company_name="Mirabelle ApS", org_number="39842851",
        country="DK", email=email, accountant_email=accountant_email,
        accountant_name="Anna Hansen", accountant_auto_send=auto_send,
        cvr_verified_source="cvrapi.dk · demo", cvr_verified_at=datetime(2026, 9, 1),
    )
    db.add(p); db.commit(); db.refresh(p)
    return p


def _real_profile(db, user, **kw):
    p = BusinessProfile(
        user_id=user.id, company_name="Mirabelle ApS", org_number="39842851",
        country="DK", email="owner@mirabelle.dk",
        accountant_email=kw.pop("accountant_email", "anna@revisor.dk"),
        accountant_name=kw.pop("accountant_name", "Anna"), **kw,
    )
    db.add(p); db.commit(); db.refresh(p)
    return p


def _row(db, user, d, rev, *, status="confirmed", demo=False):
    c = DailyClose(
        id=uuid.uuid4(), user_id=user.id, branch_id=None, date=d,
        revenue_categories=encode_breakdown({"food": rev}), revenue_total=rev,
        payment_categories=encode_breakdown({"card": rev}), payment_total=rev,
        moms_total=round(rev / 5, 2), revenue_ex_moms=round(rev - rev / 5, 2),
        moms_mode="auto", status=status, closed_by="Lars",
        closed_at=datetime.combine(d, datetime.min.time()) + timedelta(hours=21)
        if status == "confirmed" else None,
        is_deleted=False, notes=("sample · demo" if demo else "rigtig dag"),
    )
    db.add(c); db.commit(); db.refresh(c)
    return c


# ─── Must-fix 1: the seeder writes only reserved addresses, auto-send off ──


def test_a_fresh_demo_seed_has_only_reserved_addresses_and_auto_send_off(db_session):
    from app.services.demo_seed import seed_for_user
    user = _make_user(db_session, email="ny@cafe.dk")
    assert seed_for_user(db_session, user)["ok"] is not False
    p = db_session.query(BusinessProfile).filter_by(user_id=user.id).first()
    addrs = [v for v in (p.email, p.accountant_email) if v]
    assert addrs == ["info@mirabelle.example", "revisor@mirabelle.example"]
    # RFC 2606 reserved TLD — undeliverable by definition.
    assert all(a.endswith(".example") for a in addrs)
    assert p.accountant_auto_send is False
    assert p.accountant_auto_send_effective is False
    assert p.accountant_is_demo is True
    # Everything else the demo shows is still there.
    assert p.accountant_name == "Anna Hansen" and p.company_name == "Mirabelle ApS"
    assert p.cvr_verified_source.endswith(" · demo")


# ─── Must-fix 2: already-seeded accounts — the code guard ──────────────


def test_locking_a_real_close_on_a_seeded_profile_never_mails_the_demo_revisor(
        db_session, client, mailbox):
    user = _make_user(db_session)
    _demo_profile(db_session, user)  # old seed: anna@revisor.dk, auto-send NULL
    prof = client.get("/api/business", headers=_auth(user)).json()
    assert prof["accountant_is_demo"] is True
    assert prof["accountant_auto_send_effective"] is False

    r = _lock(client, user)  # a REAL close (no demo notes)
    assert r.status_code == 200, r.text
    ritual = r.json()["close_ritual"]
    assert ritual["accountant_skip_reason"] == "demo_recipient"
    assert ritual["accountant_included"] is False
    assert "anna@revisor.dk" not in ritual["sent_to"]
    assert _revisor_mails(mailbox) == []
    assert all("anna@revisor.dk" not in p["to"] for p in mailbox.sent)
    # The owner is told, in their own copy.
    owner = [p for p in mailbox.sent if p["to"] == ["anders@mirabelle.dk"]]
    assert len(owner) == 1 and DEMO_LINE_DA in owner[0]["html"]
    row = client.get(f"/api/daily-close/{r.json()['id']}", headers=_auth(user)).json()
    assert "anna@revisor.dk" not in row["email_sent_to"]
    # Recorded on the audit row.
    a = (db_session.query(AuditLog).filter(AuditLog.action == "close.auto_emailed")
         .order_by(AuditLog.created_at.desc()).first())
    assert json.loads(a.after_state)["accountant_skip_reason"] == "demo_recipient"


def test_the_guard_holds_even_with_auto_send_switched_on(db_session, client, mailbox):
    user = _make_user(db_session)
    _demo_profile(db_session, user, accountant_email="revisor@mirabelle.example", auto_send=True)
    r = _lock(client, user)
    assert r.json()["close_ritual"]["accountant_skip_reason"] == "demo_recipient"
    assert _revisor_mails(mailbox, "revisor@mirabelle.example") == []


def test_resend_on_a_seeded_profile_sends_only_the_owners_copy(db_session, client, mailbox):
    user = _make_user(db_session)
    _demo_profile(db_session, user)
    cid = _lock(client, user).json()["id"]
    mailbox.sent.clear()
    rr = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-demo-r1", "force": True}, headers=_auth(user))
    assert rr.status_code == 200, rr.text
    assert rr.json()["close_ritual"]["accountant_skip_reason"] == "demo_recipient"
    assert [p["to"] for p in mailbox.sent] == [["anders@mirabelle.dk"]]
    assert DEMO_LINE_DA in mailbox.sent[0]["html"]


def test_period_moms_payroll_and_invite_refuse_the_demo_revisor(db_session, client, mailbox):
    user = _make_user(db_session)
    _demo_profile(db_session, user)
    _row(db_session, user, date(2026, 9, 25), 1000.0)
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender, \
         patch("app.services.email_service.send_email", return_value=True) as plain:
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "pdf"}, headers=_auth(user))
        assert r.status_code == 409 and r.json()["detail"]["code"] == "demo_recipient"
        assert DEMO_LINE_DA in r.json()["detail"]["message_da"]
        # An old client echoing the seeded address gets the same answer.
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "pdf", "accountant_email": "anna@revisor.dk"},
                        headers=_auth(user))
        assert r.status_code == 409 and r.json()["detail"]["code"] == "demo_recipient"
        t = client.post("/api/tax/filing-pdf/send-to-accountant"
                        "?period_start=2026-09-01&period_end=2026-09-30",
                        json={"cc_self": True}, headers=_auth(user))
        assert t.status_code == 409 and t.json()["detail"]["code"] == "demo_recipient", t.text
        pr = client.post("/api/staff/payroll/send-to-accountant",
                         json={"period_start": "2026-09-01", "period_end": "2026-09-30"},
                         headers=_auth(user))
        assert pr.status_code == 409 and pr.json()["detail"]["code"] == "demo_recipient", pr.text
        inv = client.post("/api/accountants/invite", json={"email": "anna@revisor.dk"},
                          headers=_auth(user))
        assert inv.status_code == 409 and inv.json()["detail"]["code"] == "demo_recipient", inv.text
        assert sender.call_count == 0 and plain.call_count == 0
    assert mailbox.sent == []


def test_a_real_revisor_on_a_demo_tagged_profile_is_mailed(db_session, client, mailbox):
    """Only the SEEDED address is fenced: an owner who typed their own revisor
    onto a demo-tagged profile still has them mailed."""
    user = _make_user(db_session)
    _demo_profile(db_session, user, accountant_email="min@revisor.dk", auto_send=True)
    r = _lock(client, user)
    assert r.json()["close_ritual"]["accountant_included"] is True
    assert len(_revisor_mails(mailbox, "min@revisor.dk")) == 1


def test_the_same_address_on_a_real_profile_is_not_fenced(db_session):
    from app.services.revisor_mail import is_demo_revisor, owner_copy_allowed
    user = _make_user(db_session)
    real = _real_profile(db_session, user)  # no " · demo" tag
    assert is_demo_revisor(real) is False
    assert real.accountant_auto_send_effective is True
    assert owner_copy_allowed(real, "info@mirabelle.dk") is True


def test_the_owner_copy_never_goes_to_a_seeded_business_address(db_session, client, mailbox):
    from app.services.revisor_mail import owner_copy_allowed
    user = _make_user(db_session, email="info@mirabelle.dk")
    prof = _demo_profile(db_session, user)
    assert owner_copy_allowed(prof, "info@mirabelle.dk") is False
    assert owner_copy_allowed(prof, "info@mirabelle.example") is False
    r = _lock(client, user)
    assert r.status_code == 200, r.text
    assert r.json()["close_ritual"]["email_status"] == "skipped_no_recipient"
    assert mailbox.sent == []


def test_moms_and_period_sends_never_copy_a_seeded_business_address(db_session, client):
    """A real revisor typed onto a demo-tagged profile is mailed — but the
    owner's copy never goes to an address the demo seeder wrote."""
    user = _make_user(db_session, email="info@mirabelle.dk")
    _demo_profile(db_session, user, accountant_email="min@revisor.dk", auto_send=True)
    _row(db_session, user, date(2026, 9, 25), 1000.0)
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        t = client.post("/api/tax/filing-pdf/send-to-accountant"
                        "?period_start=2026-09-01&period_end=2026-09-30",
                        json={"cc_self": True}, headers=_auth(user))
        assert t.status_code == 200, t.text
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "pdf", "cc_self": True}, headers=_auth(user))
        assert r.status_code == 200, r.text
    assert [c.args[0] for c in sender.call_args_list] == ["min@revisor.dk", "min@revisor.dk"]


# ─── Must-fix 3: demo days stay out of the revisor's period ────────────


def _mixed_period(db, user):
    """3 real locked days + 1 real draft, 4 demo locked days + 1 demo draft."""
    real = [_row(db, user, date(2026, 9, d), 1000.0 * d) for d in (2, 3, 4)]
    _row(db, user, date(2026, 9, 5), 999.0, status="draft")
    demo = [_row(db, user, date(2026, 9, d), 20000.0, demo=True) for d in (10, 11, 12, 13)]
    _row(db, user, date(2026, 9, 14), 777.0, status="draft", demo=True)
    return real, demo


def test_a_period_mixing_demo_and_real_days_counts_only_the_real_ones(db_session, client):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    _mixed_period(db_session, user)

    rc = client.get("/api/daily-close/range-counts",
                    params={"from": "2026-09-01", "to": "2026-09-30"}, headers=_auth(user)).json()
    assert (rc["n_locked"], rc["n_drafts"], rc["n_demo"]) == (3, 1, 5)
    assert [x["date"] for x in rc["locked"]] == ["2026-09-02", "2026-09-03", "2026-09-04"]

    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "xlsx"}, headers=_auth(user))
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["n_closes"], body["n_drafts"], body["n_demo"]) == (3, 1, 5)
    (args, kwargs), _owner = sender.call_args_list
    html = args[2]
    assert "3 låste lukninger" in html
    assert "9.000,00 kr." in html          # 2.000 + 3.000 + 4.000 — no sample days
    assert "89.000,00 kr." not in html
    # The attached workbook holds the real days only.
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(kwargs["attachment_bytes"]))
    cells = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                     if c.value is not None)
    assert "2026-09-10" not in cells and "20000" not in cells
    audit = (db_session.query(AuditLog)
             .filter(AuditLog.action == "daily_close.send_to_accountant").first())
    after = json.loads(audit.after_state)
    assert (after["n_closes"], after["n_demo"], after["total_revenue"]) == (3, 5, 9000.0)


def test_the_downloadable_period_exports_leave_demo_days_out(db_session, client):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    _mixed_period(db_session, user)
    q = {"from": "2026-09-01", "to": "2026-09-30"}
    csv_txt = client.get("/api/daily-close/export.csv", params=q,
                         headers=_auth(user)).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(csv_txt), delimiter=";"))
    dates = [r[0] for r in rows[1:] if r and r[0]]
    assert dates == ["2026-09-02", "2026-09-03", "2026-09-04", "2026-09-05"]
    assert "20000" not in csv_txt and "sample · demo" not in csv_txt
    assert "I alt — 3 låste dage" in csv_txt
    pdf = pdf_text(client.get("/api/daily-close/export.pdf", params=q, headers=_auth(user)).content)
    assert "20.000,00" not in pdf and "9.000,00" in pdf
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/api/daily-close/export.xlsx", params=q,
                                             headers=_auth(user)).content))
    vals = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert 20000 not in vals and 20000.0 not in vals


def test_period_totals_never_sum_a_demo_day_handed_to_it(db_session):
    from app.services.daily_close_range_export import closes_to_csv_bytes, period_totals
    user = _make_user(db_session)
    real, demo = _mixed_period(db_session, user)
    t = period_totals(real + demo)
    assert (t["n_confirmed"], t["n_demo"], t["revenue"]) == (3, 4, 9000.0)
    assert "2026-09-10" not in closes_to_csv_bytes(real + demo).decode("utf-8-sig")


def test_a_demo_only_period_is_refused_with_nothing_locked(db_session, client):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    for d in (10, 11):
        _row(db_session, user, date(2026, 9, d), 20000.0, demo=True)
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "pdf"}, headers=_auth(user))
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == "nothing_locked"
    assert r.json()["detail"]["n_demo"] == 2
    assert sender.call_count == 0


# ─── Polish: text part, unchanged re-lock, CSP, wording ────────────────


def test_revisor_mails_carry_a_text_part(db_session, client, mailbox):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    _lock(client, user)
    rev = _revisor_mails(mailbox)[0]
    assert rev["html"] and rev["text"]
    assert "Hej Anna," in rev["text"] and "Omsætning | 12.500,00 kr." in rev["text"]
    assert "<" not in rev["text"].replace("<http", "")
    # The opt-out link survives as a readable URL.
    assert re.search(r"afmelde dem her \(https?://\S+/api/email/unsubscribe\?token=", rev["text"])
    owner = [p for p in mailbox.sent if p["to"] == ["anders@mirabelle.dk"]][0]
    assert owner["text"] and "afmelde" not in owner["text"]

    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                    json={"fmt": "pdf"}, headers=_auth(user))
    (args, kwargs), (oargs, okwargs) = sender.call_args_list
    assert "1 låst lukning" in kwargs["text"] and "afmelde dem her (http" in kwargs["text"]
    assert okwargs["text"] and "Din kopi" in okwargs["text"]


def test_an_unchanged_relock_does_not_mail_the_revisor_a_correction(db_session, client, mailbox):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    cid = _lock(client, user, rev=12500.0).json()["id"]
    assert len(_revisor_mails(mailbox)) == 1
    _unlock(client, user, cid, reason="Tjekker bare")
    r = _lock(client, user, rev=12500.0)  # nothing changed
    ritual = r.json()["close_ritual"]
    assert ritual["accountant_skip_reason"] == "unchanged"
    assert ritual["accountant_included"] is False
    assert len(_revisor_mails(mailbox)) == 1  # no "Rettet kasserapport"
    owner = [p for p in mailbox.sent if p["to"] == ["anders@mirabelle.dk"]][-1]
    assert "fik ikke en ny mail — tallene er uændrede" in owner["html"]
    # Recorded, and kept on the close so History still says "uændret".
    row = client.get(f"/api/daily-close/{cid}", headers=_auth(user)).json()
    assert row["email_status"] == "sent" and row["email_error"] == "revisor_unchanged"
    a = (db_session.query(AuditLog).filter(AuditLog.action == "close.auto_emailed")
         .order_by(AuditLog.created_at.desc()).first())
    after = json.loads(a.after_state)
    assert after["accountant_skip_reason"] == "unchanged" and after["unchanged_since"]
    # Not counted towards the cap (no mail went to the revisor).
    assert db_session.query(AuditLog).filter(
        AuditLog.action == "daily_close.revisor_lock_mail").count() == 1
    # "Send igen" is still there — an explicit send marks it as the correction.
    rr = client.post(f"/api/daily-close/{cid}/resend-email",
                     json={"key": "click-unch-1"}, headers=_auth(user))
    assert rr.status_code == 200, rr.text
    last = _revisor_mails(mailbox)[-1]
    assert len(_revisor_mails(mailbox)) == 2
    assert last["subject"].startswith("Rettet kasserapport")
    assert "Tallene er de samme" in last["html"]


def test_a_changed_relock_still_mails_the_correction(db_session, client, mailbox):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    cid = _lock(client, user, rev=12500.0).json()["id"]
    _unlock(client, user, cid)
    r = _lock(client, user, rev=12600.0)
    assert r.json()["close_ritual"]["accountant_skip_reason"] is None
    assert _revisor_mails(mailbox)[-1]["subject"].startswith("Rettet kasserapport")


def test_the_opt_out_page_renders_its_style_under_its_own_csp(db_session, unsub_db, client):
    from app.services.revisor_mail import REVISOR_TOPIC, address_fingerprint
    from app.utils.email_unsubscribe_token import make_unsubscribe_token
    user = _make_user(db_session)
    _real_profile(db_session, user)
    tok = make_unsubscribe_token(str(user.id), REVISOR_TOPIC, ttl_days=30,
                                 extra={"r": address_fingerprint("anna@revisor.dk")})
    for resp in (client.get(f"/api/email/unsubscribe?token={tok}"),
                 client.post(f"/api/email/unsubscribe?token={tok}")):
        assert resp.status_code == 200
        csp = resp.headers["content-security-policy"]
        m = re.search(r"<style>(.*?)</style>", resp.text, re.S)
        assert m, "the page carries its one stylesheet"
        digest = base64.b64encode(hashlib.sha256(m.group(1).encode("utf-8")).digest()).decode()
        assert f"style-src 'sha256-{digest}'" in csp
        assert "unsafe-inline" not in csp and "script-src" not in csp
        assert "default-src 'none'" in csp and "frame-ancestors 'none'" in csp
        assert "form-action 'self'" in csp
        # No style="" attribute is left for the hash-only policy to block.
        assert 'style="' not in resp.text
        assert ".btn-danger{" in m.group(1) or "Fortryd" in resp.text
    # Every other API response keeps the strict default.
    other = client.get("/api/daily-close", headers=_auth(user))
    assert other.headers["content-security-policy"] == "default-src 'none'; frame-ancestors 'none'"


def test_the_owners_copy_of_the_period_mail_is_not_addressed_to_the_revisor(db_session, client):
    user = _make_user(db_session)
    _real_profile(db_session, user)
    _row(db_session, user, date(2026, 9, 2), 1000.0)
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                    json={"fmt": "pdf"}, headers=_auth(user))
    (args, _k), (oargs, _ok) = sender.call_args_list
    assert "<p>Hej Anna,</p>" in args[2]
    assert "<p>Hej,</p>" in oargs[2] and "Hej Anna" not in oargs[2]


def test_a_missing_z_bon_photo_is_worded_for_who_reads_it():
    from app.routers.daily_close import _build_close_email_html
    from tests.test_revisor_artifacts import _close
    dc = _close(date(2026, 9, 25), 12500.0, 2500.0)
    common = dict(business_name="Mirabelle ApS", dc=dc, currency="DKK", closed_by="Lars",
                  has_scan=False, scan_degraded=True, is_danish=True)
    _s, rev = _build_close_email_html(**common, audience="revisor")
    _s, own = _build_close_email_html(**common, audience="owner")
    assert "Z-bon-fotoet er ikke vedhæftet — ejeren kan sende det fra BonBox." in rev
    assert "kunne ikke hentes lige nu" not in rev
    assert "tryk Send igen i Historik" in own
