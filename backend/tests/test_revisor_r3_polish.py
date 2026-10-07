"""Revisor artifacts, round 3 — the polish that reaches the revisor.

Content-reading tests (pypdf / openpyxl / the zip XML / the mail payloads) for
each polish item that changes behaviour. Every send is stubbed.
"""
from __future__ import annotations

import io
import json
import re
import uuid
import zipfile
from datetime import date, datetime, timedelta
from unittest.mock import patch

import pytest

from app.models.audit_log import AuditLog
from app.models.business_profile import BusinessProfile
from app.models.daily_close import DailyClose, encode_breakdown
from app.services.close_kasserapport_pdf import band_line_markup, build_close_kasserapport_pdf
from app.services.daily_close_range_export import (
    build_daily_close_range_pdf,
    build_daily_close_range_xlsx,
    period_document_id,
)
from tests.test_revisor_artifacts import (  # noqa: F401 — fixtures
    _auth,
    _lock,
    _make_profile,
    _make_user,
    _revisor_mails,
    _unlock,
    client,
    db_session,
    mailbox,
    pdf_text,
    unsub_db,
)


def _flat(txt: str) -> str:
    return " ".join(txt.split())


def _day(d, *, rev=1000.0, moms=200.0, diff=0.0, status="confirmed", **kw):
    base = dict(
        id=uuid.uuid4(), user_id=uuid.uuid4(), branch_id=None, date=d,
        revenue_categories=encode_breakdown({"food": rev}), revenue_total=rev,
        payment_categories=encode_breakdown({"cash": rev}), payment_total=rev,
        moms_total=moms, revenue_ex_moms=round(rev - moms, 2), moms_mode="auto",
        cash_expected=rev, cash_counted=rev + diff, cash_difference=diff,
        status=status, closed_by="Lars",
        closed_at=datetime.combine(d, datetime.min.time()) + timedelta(hours=21) if status == "confirmed" else None,
        is_deleted=False,
    )
    base.update(kw)
    return DailyClose(**base)


def _save(db, user, c):
    c.user_id = user.id
    db.add(c); db.commit(); db.refresh(c)
    return c


# ─── The kasserapport ────────────────────────────────────────────────


def test_a_short_cash_line_ends_in_one_full_stop(db_session):
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    c = _save(db_session, user, _day(date(2026, 9, 27), diff=-20.0))
    txt = pdf_text(build_close_kasserapport_pdf(db_session, user, c, profile=prof)["pdf"])
    assert "Kassen mangler 20,00 kr." in txt
    assert "kr.." not in txt


def test_a_failing_check_is_the_darkest_line_in_the_band():
    """Black-and-white print: the failing line is ink and bold with the ×
    mark; a passing line is regular weight; information is muted."""
    fail = band_line_markup({"ok": False, "text": "Kontant optalt — kassedifference -180,00 kr. …"})
    ok = band_line_markup({"ok": True, "text": "Salgsmoms beregnet af BonBox."})
    info = band_line_markup({"ok": True, "info": True, "text": "Ingen andre bilag."})
    assert "Helvetica-Bold" in fail and "#171717" in fail and "×" in fail
    assert "Helvetica-Bold" not in ok and "✓" in ok
    assert "Helvetica-Bold" not in info and info.count("·") == 1


def test_generated_time_is_on_the_venues_clock(db_session, monkeypatch):
    monkeypatch.setattr("app.utils.time.utc_now", lambda: datetime(2026, 10, 7, 21, 19))
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    c = _save(db_session, user, _day(date(2026, 9, 25)))
    txt = pdf_text(build_close_kasserapport_pdf(db_session, user, c, profile=prof)["pdf"])
    assert "Genereret 07.10.2026 kl. 23:19" in txt and "UTC" not in txt
    ptxt = pdf_text(build_daily_close_range_pdf(
        [c], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30), business_name="M"))
    assert "Genereret 07.10.2026 kl. 23:19" in ptxt and "UTC" not in ptxt


# ─── The period PDF ───────────────────────────────────────────────────


def _two_months():
    days = [_day(date(2026, 8, 1) + timedelta(days=i), rev=1000.0 + i, moms=round((1000.0 + i) / 5, 2),
                 diff=(-150.0 if i == 40 else 0.0))
            for i in range(61)]
    return days


def test_period_pdf_names_itself_on_every_page_with_a_dokument_id():
    days = _two_months()
    pdf = build_daily_close_range_pdf(
        days, from_date=date(2026, 8, 1), to_date=date(2026, 9, 30),
        business_name="Mirabelle ApS", currency="DKK", bilagsnummer="KRP-20260801-20260930")
    from pypdf import PdfReader
    pages = [p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages]
    assert len(pages) >= 3
    did = period_document_id(days, date(2026, 8, 1), date(2026, 9, 30))
    for i, page in enumerate(pages, start=1):
        flat = _flat(page)
        assert f"Dokument-id: {did}" in flat
        assert f"Side {i} af {len(pages)}" in flat
        if i > 1:
            assert ("Kasserapporter · Mirabelle ApS · 1. aug 2026 – 30. sep 2026 · "
                    f"Bilagsnr KRP-20260801-20260930 · Side {i} af {len(pages)}") in flat
    whole = "\n".join(pages)
    assert "Doc-hash" not in whole
    assert "I alt — 61 låste dage" in _flat(whole)
    assert "I alt (låste)" not in whole
    # The same closes give the same id on every download.
    assert period_document_id(list(reversed(days)), date(2026, 8, 1), date(2026, 9, 30)) == did


def test_the_kasseafstemning_heading_never_stands_alone():
    pdf = build_daily_close_range_pdf(
        _two_months(), from_date=date(2026, 8, 1), to_date=date(2026, 9, 30),
        business_name="Mirabelle ApS", currency="DKK")
    from pypdf import PdfReader
    pages = [p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages]
    page = next(p for p in pages if "Kasseafstemning" in p)
    after = page.split("Kasseafstemning", 1)[1]
    # The next line on the SAME page is a reconciliation line.
    first = next(ln for ln in after.splitlines() if ln.strip())
    assert "optalt" in first, page
    # Only the day with a difference gets its own line; the rest are counted.
    flat = _flat("\n".join(pages))
    assert "60 dage: optalt = forventet (0,00 kr.)" in flat
    assert "10. sep 2026 — optalt 890,00 kr. − forventet 1.040,00 kr. = -150,00 kr. (kassen mangler)" in flat


def test_period_pdf_marks_relocked_days_and_the_source_like_excel_and_csv():
    a, b = _day(date(2026, 9, 2)), _day(date(2026, 9, 4))
    hist = {str(b.id): ("Låst 04.09.2026 kl. 23:00 · Låst op 06.09.2026 kl. 10:00 af ejeren — "
                        "årsag: Forkert optælling · Låst igen 06.09.2026 kl. 10:12")}
    src = {str(a.id): "Z-bon (scannet) · 2 terminaler lagt sammen: 500,00 kr. + 500,00 kr.",
           str(b.id): "Indtastet af kasseansvarlig"}
    txt = _flat(pdf_text(build_daily_close_range_pdf(
        [a, b], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30), business_name="M",
        history=hist, sources=src)))
    assert "genlåst" in txt
    assert "Kilde, salgsmoms og historik pr. dag" in txt
    assert "Kilde: Z-bon (scannet) 1 dag · Indtastet af kasseansvarlig 1 dag" in txt
    assert "2. sep 2026: Z-bon (scannet) · 2 terminaler lagt sammen" in txt
    assert "4. sep 2026: Låst 04.09.2026 kl. 23:00 · Låst op 06.09.2026 kl. 10:00 af ejeren" in txt


# ─── The Excel ────────────────────────────────────────────────────────


def test_excel_carries_its_identity_print_setup_vat_column_and_cached_totals(monkeypatch):
    from openpyxl import load_workbook
    monkeypatch.setattr("app.utils.time.utc_now", lambda: datetime(2026, 10, 7, 21, 19))
    days = [_day(date(2026, 9, 24), rev=1000.0, moms=200.0),
            _day(date(2026, 9, 25), rev=2500.0, moms=500.0),
            _day(date(2026, 9, 26), rev=30053.0, moms=6010.6, status="draft")]
    f, t = date(2026, 9, 1), date(2026, 9, 30)
    raw = build_daily_close_range_xlsx(days, from_date=f, to_date=t, business_name="Mirabelle ApS",
                                       currency="DKK", bilagsnummer="KRP-20260901-20260930")
    wb = load_workbook(io.BytesIO(raw))
    s1 = wb["Oversigt"]
    kv = {s1.cell(row=r, column=1).value: s1.cell(row=r, column=2).value for r in range(1, s1.max_row + 1)}
    assert kv["Bilagsnr."] == "KRP-20260901-20260930"
    # The same id the period PDF of the same period prints.
    assert kv["Dokument-id"] == period_document_id(days, f, t)
    pdf_txt = pdf_text(build_daily_close_range_pdf(days, from_date=f, to_date=t, business_name="M"))
    assert f"Dokument-id: {kv['Dokument-id']}" in pdf_txt
    assert kv["Genereret"] == "07.10.2026 kl. 23:19"
    assert "Opbevares i 5 år efter bogføringsloven." in kv

    ws = wb["Kasserapport"]
    hdr = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    assert "Momsopgørelse" in hdr and "Bogføring" in hdr
    assert ws.cell(row=2, column=hdr.index("Momsopgørelse") + 1).value == "Salgsmoms (25 %)"
    # Print setup: landscape, one page wide, the header row on every page.
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1 and ws.sheet_properties.pageSetUpPr.fitToPage
    assert "1:1" in str(ws.print_title_rows).replace("$", "")
    # Totals: a live SUMIFS AND its value, for a previewer that never
    # recalculates — locked rows only (3.500, not 33.553).
    totals_row = 5
    assert ws.cell(row=totals_row, column=1).value == "I alt — 2 låste dage"
    assert ws.cell(row=totals_row, column=4).value.startswith("=SUMIFS(")
    # Read the way a previewer reads it — the cached value — whichever XML
    # shape the writer used (`<v/>` without lxml, `<v></v>` with it).
    cached = load_workbook(io.BytesIO(raw), data_only=True)["Kasserapport"]
    assert cached.cell(row=totals_row, column=4).value == pytest.approx(3500.0)
    assert cached.cell(row=totals_row, column=5).value == pytest.approx(700.0)


@pytest.mark.parametrize("empty", ["<v/>", "<v />", "<v></v>", ""])
def test_cached_totals_are_written_whatever_shape_the_empty_value_has(empty):
    """openpyxl writes an empty formula value as `<v />` on its own and as
    `<v></v>` when lxml is installed; both — and no <v> at all — get the total."""
    from openpyxl import load_workbook
    from app.services.daily_close_range_export import _with_cached_values
    sheet = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             '<sheetData><row r="1"><c r="A1"><f>SUM(B1:C1)</f>' + empty + '</c>'
             '<c r="B1"><v>1</v></c><c r="C1"><v>2</v></c></row></sheetData></worksheet>')
    from openpyxl import Workbook
    wb = Workbook(); wb.active["A1"] = "=SUM(B1:C1)"; wb.active["B1"] = 1; wb.active["C1"] = 2
    buf = io.BytesIO(); wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = (sheet.encode() if item.filename == "xl/worksheets/sheet1.xml"
                    else zin.read(item.filename))
            zout.writestr(item, data)
    fixed = _with_cached_values(out.getvalue(), 1, {"A1": 3.0})
    ws = load_workbook(io.BytesIO(fixed), data_only=True).active
    assert ws["A1"].value == pytest.approx(3.0)
    assert load_workbook(io.BytesIO(fixed)).active["A1"].value == "=SUM(B1:C1)"


# ─── Unlock history: the role, never the login e-mail ─────────────────


def test_unlock_history_names_the_owner_never_the_login(db_session, client, mailbox):
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    cid = _lock(client, user, rev=12500.0).json()["id"]
    _unlock(client, user, cid, reason="Forkert kortbeløb")
    dc = db_session.query(DailyClose).filter(DailyClose.id == uuid.UUID(cid)).first()
    db_session.refresh(dc)
    assert dc.unlocked_by == "owner"
    row = (db_session.query(AuditLog)
           .filter(AuditLog.action == "daily_close.unlock").order_by(AuditLog.created_at.desc()).first())
    after = json.loads(row.after_state)
    assert after["unlocked_by"] == "owner" and after["unlocked_by_email"] == user.email
    _lock(client, user, rev=13000.0)
    db_session.refresh(dc)
    txt = _flat(pdf_text(build_close_kasserapport_pdf(db_session, user, dc, profile=prof)["pdf"]))
    assert "af ejeren — årsag: Forkert kortbeløb" in txt
    assert user.email not in txt
    corr = _revisor_mails(mailbox)[-1]["html"]
    assert "Låst op af ejeren — årsag: Forkert kortbeløb" in corr
    assert user.email not in corr


# ─── The single-close PDF response ────────────────────────────────────


def test_single_pdf_is_never_cached_and_a_deleted_close_has_none(db_session, client):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    c = _save(db_session, user, _day(date(2026, 9, 25)))
    r = client.get(f"/api/daily-close/{c.id}/pdf", headers=_auth(user))
    assert r.status_code == 200
    assert r.headers["cache-control"] == "private, no-store"
    d = _save(db_session, user, _day(date(2026, 9, 26), status="draft"))
    assert client.delete(f"/api/daily-close/{d.id}", headers=_auth(user)).status_code == 204
    assert client.get(f"/api/daily-close/{d.id}/pdf", headers=_auth(user)).status_code == 404


# ─── The period send ──────────────────────────────────────────────────


def test_a_drafts_only_period_is_never_mailed(db_session, client):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    _save(db_session, user, _day(date(2026, 9, 26), status="draft"))
    with patch("app.services.email_service.send_email_with_attachment",
               return_value=(True, None)) as sender:
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "xlsx"}, headers=_auth(user))
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == "nothing_locked"
    assert sender.call_count == 0
    assert db_session.query(AuditLog).filter(
        AuditLog.action == "daily_close.send_to_accountant").count() == 0


def test_a_period_send_leaves_a_record_the_panel_reads(db_session, client):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    _save(db_session, user, _day(date(2026, 9, 25)))
    _save(db_session, user, _day(date(2026, 9, 26), status="draft"))
    with patch("app.services.email_service.send_email_with_attachment", return_value=(True, None)):
        r = client.post("/api/daily-close/send-to-accountant?from=2026-09-01&to=2026-09-30",
                        json={"fmt": "xlsx"}, headers=_auth(user))
    assert r.status_code == 200, r.text
    assert r.json()["cc_to"] == "anders@mirabelle.dk"
    rec = client.get("/api/daily-close/accountant-sends", headers=_auth(user))
    assert rec.status_code == 200
    (last,) = rec.json()
    assert last["recipient"] == "anna@revisor.dk" and last["format"] == "xlsx"
    assert (last["from"], last["to"]) == ("2026-09-01", "2026-09-30")
    assert last["n_closes"] == 1 and last["n_drafts"] == 1
    # Another tenant sees nothing of it.
    other = _make_user(db_session, email="other@cafe.dk")
    assert client.get("/api/daily-close/accountant-sends", headers=_auth(other)).json() == []


# ─── The lock mail ────────────────────────────────────────────────────


def test_owner_copy_and_revisor_reply_to_are_the_same_address(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    _lock(client, user)
    revisor = _revisor_mails(mailbox)[-1]
    owner = [p for p in mailbox.sent if p["to"] != ["anna@revisor.dk"]][-1]
    assert owner["to"] == ["anders@mirabelle.dk"]
    assert revisor["reply_to"] == "anders@mirabelle.dk" == owner["reply_to"]
    # "Afleveret til mailserveren" — never a claim of delivery.
    assert "afleveret til mailserveren til revisoren (anna@revisor.dk)" in owner["html"]


def test_an_owner_only_setup_can_resend_its_failed_copy(db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user, accountant_email=None)
    mailbox.fail = True
    cid = _lock(client, user).json()["id"]
    mailbox.fail = False
    r = client.post(f"/api/daily-close/{cid}/resend-email", json={"key": "click-own01"},
                    headers=_auth(user))
    assert r.status_code == 200, r.text
    assert r.json()["email_status"] == "sent"
    assert [p["to"] for p in mailbox.sent] == [["anders@mirabelle.dk"]]
    # A second click asks first, like a revisor send.
    again = client.post(f"/api/daily-close/{cid}/resend-email", json={"key": "click-own02"},
                        headers=_auth(user))
    assert again.status_code == 409 and again.json()["detail"]["code"] == "already_sent"


def test_a_correction_lists_each_category_once_whatever_the_key_case():
    from app.routers.daily_close import _figure_changes
    dc = _day(date(2026, 9, 25), rev=12500.0, moms=2500.0,
              revenue_categories="food:10000|drinks:2500", payment_categories="cash:2500|card:10000")
    before = {"revenue_total": 12500.0, "revenue_breakdown": {"Food": 10000.0, "Drinks": 2500.0},
              "payment_breakdown": {"Cash": 2500.0, "card": 10000.0}}
    changes, known = _figure_changes(before, dc, "DKK")
    assert known is True and changes == []
    before["revenue_breakdown"] = {"Food": 9000.0, "Drinks": 3500.0}
    changes, _ = _figure_changes(before, dc, "DKK")
    assert changes == [("Drikkevarer", "3.500,00 kr.", "2.500,00 kr."),
                       ("Mad", "9.000,00 kr.", "10.000,00 kr.")]


# ─── The revisor's opt-out ────────────────────────────────────────────


def test_the_opt_out_link_is_on_the_api_host_never_the_spa(monkeypatch):
    from app.config import settings
    from app.services.revisor_mail import revisor_unsubscribe_url
    monkeypatch.delenv("PUBLIC_API_URL", raising=False)
    monkeypatch.delenv("AIIA_REDIRECT_URI", raising=False)
    monkeypatch.setattr(settings, "FRONTEND_URL", "https://www.bonbox.dk")
    monkeypatch.setattr(settings, "PUBLIC_API_URL", "")
    monkeypatch.setattr(settings, "ENVIRONMENT", "production")
    url = revisor_unsubscribe_url(uuid.uuid4(), "anna@revisor.dk")
    assert url.startswith("https://api.bonbox.dk/api/email/unsubscribe?token=")
    # An unrelated variable no longer moves it.
    monkeypatch.setenv("AIIA_REDIRECT_URI", "https://bank.example/cb")
    assert revisor_unsubscribe_url(uuid.uuid4(), "a@b.dk").startswith("https://api.bonbox.dk/")
    # The deliberate setting does.
    monkeypatch.setattr(settings, "PUBLIC_API_URL", "https://api.staging.bonbox.dk/")
    assert revisor_unsubscribe_url(uuid.uuid4(), "a@b.dk").startswith(
        "https://api.staging.bonbox.dk/api/email/unsubscribe?token=")
    monkeypatch.setattr(settings, "PUBLIC_API_URL", "")
    monkeypatch.setattr(settings, "ENVIRONMENT", "development")
    assert revisor_unsubscribe_url(uuid.uuid4(), "a@b.dk").startswith("http://localhost:8000/")


def test_the_revisor_pages_are_danish_and_an_old_link_gets_a_danish_answer(
        db_session, unsub_db, client):
    from app.services.revisor_mail import REVISOR_TOPIC, address_fingerprint
    from app.utils.email_unsubscribe_token import make_unsubscribe_token
    user = _make_user(db_session)
    _make_profile(db_session, user)
    tok = make_unsubscribe_token(str(user.id), REVISOR_TOPIC, ttl_days=30,
                                 extra={"r": address_fingerprint("anna@revisor.dk")})
    page = client.get(f"/api/email/unsubscribe?token={tok}")
    assert page.status_code == 200 and '<html lang="da">' in page.text
    old = make_unsubscribe_token(str(user.id), REVISOR_TOPIC, ttl_days=-1,
                                 extra={"r": address_fingerprint("anna@revisor.dk")})
    for method in (client.get, client.post):
        r = method(f"/api/email/unsubscribe?token={old}")
        assert r.status_code == 410
        assert '<html lang="da">' in r.text and "Linket er udløbet" in r.text
        assert "Mirabelle ApS" in r.text and "hello@bonbox.dk" in r.text
        assert "Daily Brief" not in r.text
    # Nothing changed on the profile.
    db_session.expire_all()
    assert db_session.query(BusinessProfile).first().accountant_opted_out_at is None
    # A forged token still gets the generic page, never the revisor's.
    forged = client.get("/api/email/unsubscribe?token=abc.def")
    assert forged.status_code == 410 and "Linket er udløbet" not in forged.text
