"""Revisor artifacts, round 3 — the three things that contradicted themselves.

Each test READS the generated artifacts (pypdf for the PDFs, openpyxl for the
workbook, the csv module, the mail HTML) for the SAME days, and asserts that
the kasserapport, the period PDF, the Excel, the CSV and the lock mail say the
same thing about:

  1. readiness — ONE rule for revenue without a category: a NOTE ("Ikke
     fordelt på kategori: X kr.") whether part of the day is unsplit or all of
     it. It used to be the other way round: a partial split failed, no split at
     all passed. Whether the MOMS is right is the MOMS line's question, asked
     whatever the split (categories carry no VAT rate): 25 % within rounding
     or a whole momsfri day passes; anything else is GENNEMGÅS, with the
     deviation in kroner.
  2. the MOMS label — "Salgsmoms (25 %)" only over a figure that IS 25 % of the
     base; a figure read off the Z-bon or typed is named by its source, with
     the rate it works out to when that is not 25 %.
  3. the bilag number — every period row carries the day's own kasserapport
     number (KR-…) and its Dokument-id, so a row can be matched to its voucher.

No mail leaves the process (the Resend client is stubbed).
"""
from __future__ import annotations

import base64
import csv
import io
import json
import re
import uuid
from datetime import date, datetime

import pytest

from app.models.daily_close import DailyClose, encode_breakdown
from app.services.close_kasserapport_pdf import build_close_kasserapport_pdf
from app.services.daily_close_range_export import (
    build_daily_close_range_pdf,
    build_daily_close_range_xlsx,
    closes_to_csv_bytes,
)
from app.services.kasserapport_claims import build_close_claims, close_readiness, moms_label
from tests.test_revisor_artifacts import (  # noqa: F401 — fixtures
    _auth,
    _lock,
    _make_profile,
    _make_user,
    _revisor_mails,
    client,
    db_session,
    mailbox,
    pdf_text,
)


def _flat(txt: str) -> str:
    return " ".join(txt.split())


def _day(d, *, rev=15000.0, moms=3000.0, cats=None, mode="auto", kind=None, **kw):
    """A locked close: cash and payments tie out, so only the split and the
    MOMS decide the verdict."""
    base = dict(
        id=uuid.uuid4(), user_id=uuid.uuid4(), branch_id=None, date=d,
        revenue_categories=encode_breakdown(cats) if cats else None,
        revenue_total=rev,
        payment_categories=encode_breakdown({"cash": 5000.0, "card": rev - 5000.0}),
        payment_total=rev, moms_total=moms,
        revenue_ex_moms=(round(rev - moms, 2) if moms is not None else None),
        moms_mode=mode, cash_expected=5000.0, cash_counted=5000.0, cash_difference=0.0,
        status="confirmed", closed_by="Lars", closed_at=datetime(2026, 9, 25, 21, 0),
        source_meta=json.dumps({"kind": kind}) if kind else None, is_deleted=False,
    )
    base.update(kw)
    return DailyClose(**base)


# ─── 1. One readiness rule (claims level) ────────────────────────────


@pytest.mark.parametrize("cats", [
    None,                                   # total-only Z-bon
    {"food": 1000.0},                       # a partial split
    {"food": 14999.0},                      # 1 kr. left over
    {"food": 10000.0, "drinks": 5000.0},    # a full split
])
def test_single_rate_day_is_book_ready_with_any_split(cats):
    """Typing more categories can never turn a book-ready day into GENNEMGÅS."""
    claims = build_close_claims(_day(date(2026, 9, 24), cats=cats))
    a = claims["assurance"]
    assert a["all_ok"] is True and a["heading"] == "KLAR TIL BOGFØRING"
    lines = next(c for c in a["checks"] if c["check"] == "lines")
    if cats and sum(cats.values()) == 15000.0:
        assert lines["ok"] is True and not lines.get("info")
    else:
        unsplit = 15000.0 - sum((cats or {}).values())
        assert lines.get("info") is True
        assert lines["text"].startswith("Ikke fordelt på kategori: ")
        assert claims["unallocated_value"] == pytest.approx(unsplit)


@pytest.mark.parametrize("cats", [
    None,                                   # total-only Z-bon
    {"food": 1000.0},                       # a partial split
    {"food": 10000.0, "drinks": 5000.0},    # a full split
    {"food": 14999.0, "drinks": 1.0},       # a full split typed to clear a flag
])
def test_an_off_rate_moms_is_review_whatever_the_split(cats):
    """2.500 kr. MOMS on 15.000 kr. is 20 % of the base. Revenue categories
    carry no VAT rate, so splitting the day verifies nothing about it: the
    MOMS line itself fails — with the deviation in kroner — whatever the split.
    (A full split used to certify it KLAR, 500 kr. under-declared if every
    sale was 25 %.) The unsplit amount is a note and never claims a rate."""
    dc = _day(date(2026, 9, 26), moms=2500.0, mode="manual", kind="zbon", cats=cats)
    a = build_close_claims(dc)["assurance"]
    assert a["all_ok"] is False and a["heading"] == "GENNEMGÅS"
    assert close_readiness(dc)["ready"] is False
    moms = next(c for c in a["checks"] if c["check"] == "moms")
    assert moms["ok"] is False and not moms.get("info")
    assert moms["text"] == ("Salgsmoms aflæst fra Z-bon svarer til 20 % af omsætningen ekskl. "
                            "moms — 500,00 kr. under 25 % moms på hele omsætningen. "
                            "Kontrollér den mod Z-bonnen.")
    lines = next(c for c in a["checks"] if c["check"] == "lines")
    assert lines["ok"] is True
    for c in a["checks"]:
        assert "mere end én momssats" not in c["text"]
    if cats is None or sum(cats.values()) < 15000.0:
        assert lines.get("info") is True
        assert lines["text"].startswith("Ikke fordelt på kategori: ")
        assert "momssats" not in lines["text"]


@pytest.mark.parametrize("kind", [
    "zbon",   # a scanned Z-bon MOMS the form kept under Auto ("the scanned figure still WINS")
    None,     # BonBox's own MOMS on revenue − momsfri sales (exempt_sales_total is not stored)
])
@pytest.mark.parametrize("cats", [None, {"food": 6000.0, "drinks": 4000.0}])
def test_an_off_rate_auto_moms_is_review_too(kind, cats):
    """Pinned deliberately (removal audit, AUTO source). 1.800 kr. on 10.000 kr.
    saved under moms_mode=auto is 21,9 % of the base — not BonBox's revenue ×
    25/125 (2.000 kr.). Before the rate rule it read KLAR with "Salgsmoms
    beregnet af BonBox ud fra omsætningen", which is false for both real
    sources (read off the Z-bon; or worked out on a base net of momsfri sales
    the page does not carry). The rate rule asks the figure, not the toggle:
    the same day is GENNEMGÅS whatever the split, exactly as a Z-bon figure
    saved under Manual, and every artifact reads that one verdict."""
    dc = _day(date(2026, 9, 26), rev=10000.0, moms=1800.0, mode="auto", kind=kind, cats=cats,
              payment_categories=encode_breakdown({"cash": 5000.0, "card": 5000.0}),
              payment_total=10000.0)
    claims = build_close_claims(dc)
    a = claims["assurance"]
    assert a["all_ok"] is False and a["heading"] == "GENNEMGÅS"
    assert close_readiness(dc)["ready"] is False
    moms = next(c for c in a["checks"] if c["check"] == "moms")
    assert moms["ok"] is False and not moms.get("info")
    assert moms["text"] == ("Salgsmoms svarer til 21,9 % af omsætningen ekskl. moms — "
                            "200,00 kr. under 25 % moms på hele omsætningen. "
                            "Kontrollér den mod Z-bonnen.")
    assert claims["moms_label"] == "Salgsmoms (svarer til 21,9 %)"
    assert not any("beregnet af BonBox" in c["text"] for c in a["checks"])
    # Only the MOMS line fails: cash, payments and the split say what they say.
    assert [c["check"] for c in a["checks"] if not c["ok"]] == ["moms"]
    # BonBox's own figure on the whole revenue is still certified as such.
    std = _day(date(2026, 9, 26), rev=10000.0, moms=2000.0, mode="auto", kind=kind, cats=cats,
               payment_categories=encode_breakdown({"cash": 5000.0, "card": 5000.0}),
               payment_total=10000.0)
    assert close_readiness(std)["ready"] is True


@pytest.mark.parametrize("rev, moms, kind, label", [
    (15000.0, 2998.0, "typed", "Salgsmoms (indtastet)"),        # 2,50 kr. off
    (50000.0, 10001.20, "zbon", "Salgsmoms (fra Z-bon)"),        # 1,50 kr. off
    (28469.0, 5695.30, "zbon", "Salgsmoms (fra Z-bon)"),         # 1,87 kr. off
])
def test_a_few_kroner_of_rounding_is_the_standard_rate(rev, moms, kind, label):
    """A Z-bon's per-line VAT rounding is not another rate: the figure prints
    as 25 %, so it IS 25 % — never "svarer til 25 %, ikke 25 %" and never
    GENNEMGÅS on a total-only day."""
    dc = _day(date(2026, 9, 24), rev=rev, moms=moms, mode="manual", kind=kind)
    claims = build_close_claims(dc)
    a = claims["assurance"]
    assert claims["moms_label"] == label
    assert a["all_ok"] is True and a["heading"] == "KLAR TIL BOGFØRING"
    blob = " ".join(c["text"] for c in a["checks"]) + " " + (claims["moms_manual_note"] or "")
    assert "ikke 25 %" not in blob and "momssats" not in blob.replace("(én momssats)", "")
    moms_line = next(c for c in a["checks"] if c["check"] == "moms")
    assert moms_line["ok"] is True and "svarer til 25 %" in moms_line["text"]


def test_a_deviation_beyond_rounding_never_prints_the_same_rate_twice():
    """Just outside the tolerance on a small day: the effective rate gets a
    second decimal rather than read "25 %, ikke 25 %"."""
    from app.services.kasserapport_claims import eff_pct_text, moms_rate_info
    dc = _day(date(2026, 9, 24), rev=1000.0, moms=201.30, mode="manual", kind="typed")
    info = moms_rate_info(dc)
    assert info["off"] is True
    assert eff_pct_text(info) != "25 %"
    assert "25 %, ikke 25 %" not in moms_label(dc)


def test_a_momsfri_day_is_one_rate_and_book_ready():
    """0,00 kr. MOMS on 10.000 kr. with no categories: the whole day momsfri —
    ONE rate (0 %). Stated on a muted line, never "mere end én momssats" and
    never GENNEMGÅS (it was KLAR before the rate rule, and it is again)."""
    dc = _day(date(2026, 9, 24), rev=10000.0, moms=0.0, mode="manual", kind="typed",
              payment_categories=encode_breakdown({"cash": 5000.0, "card": 5000.0}),
              payment_total=10000.0)
    claims = build_close_claims(dc)
    a = claims["assurance"]
    assert a["all_ok"] is True and a["heading"] == "KLAR TIL BOGFØRING"
    moms_line = next(c for c in a["checks"] if c["check"] == "moms")
    assert moms_line["ok"] is True and moms_line.get("info") is True
    assert moms_line["text"] == ("Salgsmoms indtastet af kasseansvarlig: 0,00 kr. — hele "
                                 "dagens omsætning er opgjort som momsfri (0 %).")
    lines = next(c for c in a["checks"] if c["check"] == "lines")
    assert lines.get("info") is True and "mere end én" not in lines["text"]
    assert claims["moms_label"] == "Salgsmoms (indtastet, svarer til 0 %)"


def test_lines_that_contradict_the_total_still_fail():
    dc = _day(date(2026, 9, 24), cats={"food": 16000.0})
    a = build_close_claims(dc)["assurance"]
    assert a["all_ok"] is False


# ─── 2. One MOMS label rule (claims level) ───────────────────────────


@pytest.mark.parametrize("kw, label", [
    (dict(), "Salgsmoms (25 %)"),
    (dict(mode="manual", kind="zbon"), "Salgsmoms (fra Z-bon)"),
    (dict(mode="manual", kind="typed"), "Salgsmoms (indtastet)"),
    (dict(moms=2500.0, mode="manual", kind="zbon"), "Salgsmoms (fra Z-bon, svarer til 20 %)"),
    (dict(moms=2500.0, mode="manual", kind="typed"), "Salgsmoms (indtastet, svarer til 20 %)"),
    (dict(moms=None), "Salgsmoms"),
])
def test_moms_label_names_the_source_and_never_a_rate_the_figure_is_not(kw, label):
    assert moms_label(_day(date(2026, 9, 24), **kw)) == label


# ─── 1–3 together: the same days through every artifact ───────────────


def _period(db, user, branch=None):
    """Four locked days, one of each case, saved for this user."""
    days = [
        # A: total-only, standard rate → KLAR, "Salgsmoms (25 %)"
        _day(date(2026, 9, 24)),
        # B: partial split, standard rate → KLAR (a note, not a failure)
        _day(date(2026, 9, 25), cats={"food": 4000.0}),
        # C: total-only, Z-bon MOMS at 20 % → GENNEMGÅS
        _day(date(2026, 9, 26), moms=2500.0, mode="manual", kind="zbon"),
        # D: as C but fully split → still GENNEMGÅS (a split verifies no rate)
        _day(date(2026, 9, 27), moms=2500.0, mode="manual", kind="zbon",
             cats={"food": 10000.0, "drinks": 5000.0}),
    ]
    for c in days:
        c.user_id = user.id
        db.add(c)
    db.commit()
    for c in days:
        db.refresh(c)
    return days


EXPECT = {
    date(2026, 9, 24): (True, "Salgsmoms (25 %)"),
    date(2026, 9, 25): (True, "Salgsmoms (25 %)"),
    date(2026, 9, 26): (False, "Salgsmoms (fra Z-bon, svarer til 20 %)"),
    date(2026, 9, 27): (False, "Salgsmoms (fra Z-bon, svarer til 20 %)"),
}


def test_the_same_day_agrees_across_kasserapport_period_pdf_excel_csv_and_mail(db_session):
    from openpyxl import load_workbook
    from app.routers.daily_close import _build_close_email_html, _range_extras

    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    days = _period(db_session, user)
    f, t = date(2026, 9, 1), date(2026, 9, 30)

    # ── the single kasserapport of each day: verdict, label, bilag, id ──
    per_day = {}
    for c in days:
        out = build_close_kasserapport_pdf(db_session, user, c, profile=prof)
        txt = _flat(pdf_text(out["pdf"]))
        ready, label = EXPECT[c.date]
        assert ("KLAR TIL BOGFØRING" in txt) is ready, (c.date, txt)
        assert ("GENNEMGÅS" in txt) is (not ready), (c.date, txt)
        assert f"{label} " in txt, (c.date, label, txt)
        if label != "Salgsmoms (25 %)":
            # Never "25 %" over a figure that is not 25 % of the base.
            assert "Salgsmoms (25 %)" not in txt
        assert out["bilagsnummer"] == f"KR-{c.date:%Y%m%d}-{c.date:%Y%m%d}"
        assert f"Bilagsnr. {out['bilagsnummer']}" in txt
        assert f"Dokument-id: {out['doc_id']}" in txt
        if c.revenue_categories is None or c.date == date(2026, 9, 25):
            assert "Ikke fordelt på kategori" in txt
        per_day[c.date] = out

    extras = _range_extras(db_session, user, days)

    # ── the period PDF: the KR number and Dokument-id on every row ──
    period_pdf = build_daily_close_range_pdf(
        days, from_date=f, to_date=t, business_name="Mirabelle ApS", currency="DKK",
        profile=prof, db=db_session, user_id=user.id, bilagsnummer="KRP-20260901-20260930",
        tz=extras["tz"], branch_names=extras["branch_names"],
        history=extras["history"], sources=extras["sources"])
    ptxt = pdf_text(period_pdf)
    pflat = _flat(ptxt)
    for d, out in per_day.items():
        assert out["bilagsnummer"] in ptxt
        assert f"id {out['doc_id']}" in ptxt
    # One rule for the badge: the two off-rate days are the ones named.
    assert ("2 af 4 klar til bogføring · 2 skal gennemgås: 26. sep 2026, 27. sep 2026"
            in pflat)
    # The 20 % label for the Z-bon days, under the amount and in the source block.
    assert "26. sep 2026: Salgsmoms (fra Z-bon, svarer til 20 %)" in pflat
    assert "27. sep 2026: Salgsmoms (fra Z-bon, svarer til 20 %)" in pflat
    # The column header does not claim 25 % for a period that holds a 20 % day.
    assert "Moms 25%" not in ptxt

    # ── the Excel and the CSV: same KR, id, label and verdict per row ──
    wb = load_workbook(io.BytesIO(build_daily_close_range_xlsx(
        days, from_date=f, to_date=t, business_name="Mirabelle ApS", currency="DKK",
        profile=prof, db=db_session, user_id=user.id, tz=extras["tz"],
        history=extras["history"], sources=extras["sources"],
        branch_names=extras["branch_names"])))
    ws = wb["Kasserapport"]
    hdr = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    assert hdr[:3] == ["Dato", "Bilagsnr.", "Dokument-id"]
    xrows = {}
    for r in range(2, 6):
        row = dict(zip(hdr, [ws.cell(row=r, column=i).value for i in range(1, ws.max_column + 1)]))
        xrows[row["Dato"].date()] = row

    rows = list(csv.reader(io.StringIO(
        closes_to_csv_bytes(days, currency="DKK", **extras).decode("utf-8-sig")), delimiter=";"))
    chdr = rows[0]
    crows = {r[0]: dict(zip(chdr, r)) for r in rows[1:-1]}

    for d, out in per_day.items():
        ready, label = EXPECT[d]
        verdict = "Klar til bogføring" if ready else "Gennemgås: "
        for row in (xrows[d], crows[d.isoformat()]):
            assert row["Bilagsnr."] == out["bilagsnummer"]
            assert row["Dokument-id"] == out["doc_id"]
            assert row["Momsopgørelse"] == label
            assert row["Bogføring"].startswith(verdict), row["Bogføring"]
    for d in ("2026-09-26", "2026-09-27"):
        assert ("500,00 kr. under 25 % moms på hele omsætningen. Kontrollér den mod Z-bonnen"
                in crows[d]["Bogføring"])
        assert "mere end én momssats" not in crows[d]["Bogføring"]

    # ── the lock mail for each day: the same verdict, label and bilag ──
    for c in days:
        out = per_day[c.date]
        ready, label = EXPECT[c.date]
        _subj, html = _build_close_email_html(
            business_name="Mirabelle ApS", dc=c, currency="DKK", closed_by="Lars",
            has_scan=False, scan_degraded=False, is_danish=True,
            attachment_name=out["filename"], audience="revisor",
            bilagsnummer=out["bilagsnummer"], doc_id=out["doc_id"])
        plain = _flat(re.sub(r"<[^>]+>", " ", html))
        assert label in plain
        assert ("Klar til bogføring" in plain) is ready
        assert ("Gennemgås" in plain) is (not ready)
        assert f"bilagsnr. {out['bilagsnummer']}, dokument-id {out['doc_id']}" in plain
        if label != "Salgsmoms (25 %)":
            assert "(25 %)" not in plain


def test_the_lock_mail_sent_on_lock_carries_the_kasserapports_bilag_and_id(
        db_session, client, mailbox):
    user = _make_user(db_session)
    _make_profile(db_session, user)
    r = _lock(client, user, d="2026-09-24", rev=15000.0, moms_mode="manual",
              moms_total=2500.0, revenue_breakdown={},
              revenue_total_override=15000.0, revenue_total_owner_set=True)
    assert r.status_code == 200, r.text
    mail = _revisor_mails(mailbox)[-1]
    pdf = pdf_text(base64.b64decode(mail["attachments"][0]["content"]))
    doc_id = re.search(r"Dokument-id: ([0-9a-f]{16})", pdf).group(1)
    plain = _flat(re.sub(r"<[^>]+>", " ", mail["html"]))
    assert f"bilagsnr. KR-20260924-20260924, dokument-id {doc_id}" in plain
    # The figure was typed and is 20 % of the base: named by its source in
    # both the attachment and the mail body, never "(25 %)".
    assert "Salgsmoms (indtastet, svarer til 20 %)" in _flat(pdf)
    assert "Salgsmoms (indtastet, svarer til 20 %)" in plain
    assert "(25 %)" not in plain
    # Total-only and mixed-rate: GENNEMGÅS on the PDF AND in the mail.
    assert "GENNEMGÅS" in pdf and "Gennemgås" in plain


def test_period_mail_uses_the_same_label_rule_and_verdict():
    from app.routers.daily_close import _accountant_email_body
    from app.services.daily_close_range_export import period_totals
    days = [_day(date(2026, 9, 24)), _day(date(2026, 9, 26), moms=2500.0, mode="manual", kind="zbon")]
    html = _accountant_email_body(
        business_name="Mirabelle ApS", from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        totals=period_totals(days), currency="DKK", fmt="xlsx", message=None, is_danish=True)
    plain = _flat(re.sub(r"<[^>]+>", " ", html))
    assert "Salgsmoms (25 %)" not in plain and "Salgsmoms 5.500,00 kr." in plain
    assert "1 af 2 klar til bogføring · 1 skal gennemgås: 26. sep." in plain
    std = _accountant_email_body(
        business_name="Mirabelle ApS", from_date=date(2026, 9, 1), to_date=date(2026, 9, 30),
        totals=period_totals(days[:1]), currency="DKK", fmt="xlsx", message=None, is_danish=True)
    assert "Salgsmoms (25 %)" in std and "1 af 1 klar til bogføring" in std


def test_two_branches_keep_their_own_kr_number_in_every_period_row(db_session):
    from app.models.branch import Branch
    from app.routers.daily_close import _range_extras
    user = _make_user(db_session)
    prof = _make_profile(db_session, user)
    ves = Branch(id=uuid.uuid4(), user_id=user.id, name="Vesterbro")
    db_session.add(ves); db_session.commit()
    c = _day(date(2026, 9, 24), branch_id=ves.id)
    c.user_id = user.id
    db_session.add(c); db_session.commit(); db_session.refresh(c)
    kr = build_close_kasserapport_pdf(db_session, user, c, profile=prof)["bilagsnummer"]
    assert kr.startswith("KR-20260924-20260924-VES")
    extras = _range_extras(db_session, user, [c])
    csv_txt = closes_to_csv_bytes([c], currency="DKK", **extras).decode("utf-8-sig")
    assert kr in csv_txt
    ptxt = pdf_text(build_daily_close_range_pdf(
        [c], from_date=date(2026, 9, 1), to_date=date(2026, 9, 30), business_name="M",
        currency="DKK", branch_names=extras["branch_names"]))
    # A branch number wraps after a hyphen in the 30 mm column.
    assert kr in ptxt.replace("-\n", "-")
